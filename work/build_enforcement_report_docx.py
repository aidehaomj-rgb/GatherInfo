"""Render the 30-day enforcement report as a polished Word document."""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "18324A"
TEAL = "0F766E"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "E8F5F2"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
BORDER = "D0D5DD"
BLACK = "101828"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 100, "bottom": 100, "start": 120, "end": 120}


def _set_run_font(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color=BORDER, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _set_table_geometry(table, widths):
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {TABLE_WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_paragraph_border_bottom(paragraph, color=BLUE, size="12", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def _add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    _set_run_font(run, size=9, color=MID_GRAY)


def _add_hyperlink(paragraph, text, url, size=9.5):
    if not url:
        run = paragraph.add_run("公开来源链接未保存")
        _set_run_font(run, size=size, color=MID_GRAY)
        return
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.extend([r_fonts, color, underline, sz])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _add_label_paragraph(doc, label, text, after=4, size=10.5):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.1
    label_run = paragraph.add_run(label)
    _set_run_font(label_run, size=size, bold=True, color=DARK_BLUE)
    value_run = paragraph.add_run(text)
    _set_run_font(value_run, size=size)
    return paragraph


def _add_body(doc, text, after=6, size=11, color=BLACK, bold=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, color=color, bold=bold)
    return paragraph


def _add_callout(doc, label, text, fill=LIGHT_TEAL, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    _set_table_geometry(table, [TABLE_WIDTH_DXA])
    _set_table_borders(table, color=accent, size="6")
    cell = table.cell(0, 0)
    _set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(label)
    _set_run_font(run, size=10.5, bold=True, color=accent)
    run = paragraph.add_run(text)
    _set_run_font(run, size=10.5, color=BLACK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def _configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [
        ("Title", 23, BLACK, 0, 4),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("执法信息采集 | 近30日综合分析")
    _set_run_font(run, size=9, color=MID_GRAY, bold=True)
    run = paragraph.add_run(" " * 20 + "内部工作参考")
    _set_run_font(run, size=9, color=MID_GRAY)
    _set_paragraph_border_bottom(paragraph, color=BORDER, size="5", space="2")
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("GatherInfo | 第 ")
    _set_run_font(run, size=9, color=MID_GRAY)
    _add_page_field(paragraph)
    run = paragraph.add_run(" 页")
    _set_run_font(run, size=9, color=MID_GRAY)


def _add_title_block(doc, data):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run("执法信息采集")
    _set_run_font(run, size=12, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("近30日综合分析报告")
    _set_run_font(run, size=23, bold=True, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run(data["subtitle"])
    _set_run_font(run, size=13, color=MID_GRAY)
    for label, value in [
        ("统计时段：", data["date_range"] + "（北京时间口径）"),
        ("数据范围：", f"系统内发布日期可核验的{data['stats']['total']}条境外执法信息"),
        ("报告用途：", "执法动态研判、涉华线索发现及后续核查参考"),
    ]:
        _add_label_paragraph(doc, label, value, after=2, size=10.5)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    _set_paragraph_border_bottom(rule, color=BLUE, size="14", space="2")


def _add_metric_strip(doc, data):
    levels = data["stats"]["levels"]
    metrics = [
        (str(data["stats"]["total"]), "近30日案例"),
        (str(levels.get("strong", 0)), "强涉华"),
        (str(levels.get("weak", 0)), "弱涉华"),
        (str(levels.get("major_non_china", 0)), "重大非涉华"),
    ]
    table = doc.add_table(rows=1, cols=4)
    _set_table_geometry(table, [2340, 2340, 2340, 2340])
    _set_table_borders(table, color=BORDER, size="5")
    for index, (value, label) in enumerate(metrics):
        cell = table.cell(0, index)
        _set_cell_shading(cell, LIGHT_BLUE if index == 0 else LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(value)
        _set_run_font(run, size=20, bold=True, color=BLUE if index < 3 else NAVY)
        paragraph = cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        _set_run_font(run, size=9.5, color=MID_GRAY, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_category_table(doc, data):
    headings = ["分类", "数量", "强涉华", "研判重点"]
    focus = {
        "毒品及药品": "大宗毒品、旅客行李、邮包及中国来源未批准药品",
        "贸易合规与一般走私": "出口管制、原产地换标、误报、非法进口和贵金属",
        "濒危物种及野生动植物": "活体龟、蜥蜴、鸟类及CITES许可",
        "烟草及新型烟草制品": "货柜误报、冷库板夹藏、私烟及电子烟",
        "知识产权侵权": "体育赛事周边、箱包及国际邮件快件",
        "枪爆及武器": "零部件采购、非法制造及港口查获方式",
        "其他边境执法": "与进出口关联偏弱，建议降级观察",
    }
    strong_counts = defaultdict(int)
    for case in data["strong_cases"]:
        strong_counts[case["category"]] += 1
    rows = [[
        category,
        str(data["stats"]["categories"].get(category, 0)),
        str(strong_counts.get(category, 0)),
        focus[category],
    ] for category in data["category_order"]]
    table = doc.add_table(rows=1, cols=4)
    for index, heading in enumerate(headings):
        cell = table.cell(0, index)
        _set_cell_shading(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(heading)
        _set_run_font(run, size=9.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            if row_index % 2:
                _set_cell_shading(cells[col_index], "F8FAFC")
            paragraph = cells[col_index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            _set_run_font(run, size=9.3, bold=col_index == 0, color=DARK_BLUE if col_index == 0 else BLACK)
    _set_table_geometry(table, [2250, 850, 950, 5310])
    _set_table_borders(table)


def _add_strong_summary_table(doc, data):
    table = doc.add_table(rows=1, cols=4)
    headers = ["序号", "日期/地区", "分类", "核心线索"]
    for index, heading in enumerate(headers):
        cell = table.cell(0, index)
        _set_cell_shading(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(heading)
        _set_run_font(run, size=9, bold=True, color=WHITE)
    for index, case in enumerate(data["strong_cases"], 1):
        cells = table.add_row().cells
        values = [
            str(index),
            f"{case['published_date']}\n{case['jurisdiction']}",
            case["category"],
            case["detail"]["headline"],
        ]
        for col_index, value in enumerate(values):
            if index % 2 == 0:
                _set_cell_shading(cells[col_index], "F8FAFC")
            paragraph = cells[col_index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index < 3 else WD_ALIGN_PARAGRAPH.LEFT
            for line_index, line in enumerate(value.split("\n")):
                if line_index:
                    paragraph.add_run().add_break()
                run = paragraph.add_run(line)
                _set_run_font(run, size=8.8, bold=col_index == 0, color=DARK_BLUE if col_index == 3 else BLACK)
    _set_table_geometry(table, [620, 1550, 1720, 5470])
    _set_table_borders(table)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def _add_numbered_recommendation(doc, number, title, body):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(title + "。")
    _set_run_font(run, size=11, bold=True, color=DARK_BLUE)
    run = paragraph.add_run(body)
    _set_run_font(run, size=11)


def build(data, output_path):
    doc = Document()
    _configure_styles(doc)
    _configure_page(doc)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = "境外进出口执法动态与强涉华线索分析"
    doc.core_properties.author = "GatherInfo"
    doc.core_properties.keywords = "执法信息, 涉华线索, 海关风险, 境外执法"

    _add_title_block(doc, data)
    _add_metric_strip(doc, data)
    _add_callout(
        doc,
        "核心判断：",
        "强涉华风险已由单一的中国来源货物，扩展到中国籍人员、中国企业、香港和日本转运节点、跨境电商及循环贸易网络。公开源信息能够发现风险，但多数案件不会披露箱号、提单号等关键单证，必须通过海关数据和执法协作继续补链。",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    _add_body(
        doc,
        "编制说明：本报告以系统采集正文、中文译文及原始网页为依据。人名、企业、单证和运输信息只摘录公开来源明确披露的内容；未披露字段统一标注，不作推测性补全。",
        size=9.5,
        color=MID_GRAY,
    )

    doc.add_heading("一、总体态势", level=1)
    levels = data["stats"]["levels"]
    _add_body(
        doc,
        f"近30日共纳入{data['stats']['total']}条境外执法信息，其中强涉华{levels.get('strong', 0)}条、弱涉华{levels.get('weak', 0)}条、重大非涉华{levels.get('major_non_china', 0)}条。强涉华案例覆盖美国、中国台湾、泰国、菲律宾、斯里兰卡、中国香港、葡萄牙、印度尼西亚和新加坡，来源结构较以往更为多元。",
    )
    _add_category_table(doc, data)

    doc.add_heading("二、分类研判", level=1)
    category_analysis = [
        ("（一）毒品及药品", "大宗毒品案件继续呈现跨国组织化、藏匿载体机械化和运输路径多段化特征；强涉华案例同时出现中国籍人员参与、从中国采购芬太尼及合成大麻素、中国来源未批准处方药经跨境电商和邮包进入境外市场等情形。"),
        ("（二）贸易合规与一般走私", "重点风险包括先进计算设备绕道转运、自由区换标、冷冻农产品和电子烟误报、二手手机非法进口、贵金属夹藏及循环贸易骗税。后续处置必须以舱单、提单、报关单和资金流数据补齐主体链条。"),
        ("（三）濒危物种及野生动植物", "案件主要涉及活体龟、蜥蜴和鸟类等，运输渠道包括旅客行李和国际邮包。美国龟类案件已披露中国籍人员姓名、包裹数量和虚假申报品名，具备进一步串并收件人、交易平台和CITES许可记录的条件。"),
        ("（四）烟草及新型烟草制品", "风险形态包括海运货柜误报电子烟、冷库板夹藏香烟、原产地换标和一般私烟运输。应同步关注货柜级申报异常以及品牌、包装、税标等商品特征。"),
        ("（五）知识产权侵权", "美国和葡萄牙案件显示中国来源商品在体育赛事周边、箱包等品类中仍具较高侵权风险。当前公开通报多缺少生产商和进口商信息，需通过扣留案卷和包裹级数据识别国内生产、出口主体。"),
        ("（六）枪爆及其他边境执法", "枪爆案件均属重大非涉华案例，其零部件采购、非法制造和港口查获方式可补充风险规则。3条其他边境执法中部分内容与进出口关联偏弱，建议后续降级至观察池。"),
    ]
    for heading, body in category_analysis:
        doc.add_heading(heading, level=2)
        _add_body(doc, body)

    doc.add_heading("三、强涉华案例深度分析", level=1)
    _add_body(
        doc,
        "以下16条均已核验原文发布日期和涉华依据。先列总览，再逐案提取可用于后续核查的人员、企业、单证、运输和货物信息。",
    )
    _add_strong_summary_table(doc, data)
    for index, case in enumerate(data["strong_cases"], 1):
        detail = case["detail"]
        research = case["research"]
        heading = doc.add_heading(f"{index}. {detail['headline']}", level=2)
        heading.paragraph_format.keep_with_next = True
        _add_label_paragraph(doc, "核心事实：", detail["core_fact"], after=7)
        _add_label_paragraph(doc, "人名及人员线索：", detail["persons"])
        _add_label_paragraph(doc, "企业及机构线索：", detail["companies"])
        _add_label_paragraph(doc, "集装箱号、提单号等单证：", detail["documents"])
        _add_label_paragraph(doc, "运输工具及路线：", detail["transport"])
        _add_label_paragraph(doc, "涉案货物：", detail["goods"], after=7)
        _add_callout(doc, "下一步处置提示：", detail["disposition"])
        _add_label_paragraph(doc, "多语种复核范围：", research["languages"])
        _add_label_paragraph(doc, "关联报道补充：", research["findings"])
        _add_label_paragraph(doc, "图片及多模态核验：", research["multimodal"])
        _add_label_paragraph(doc, "证据判断：", research["assessment"], after=6)
        supplemental = doc.add_paragraph()
        supplemental.paragraph_format.space_after = Pt(5)
        run = supplemental.add_run("补充来源：")
        _set_run_font(run, size=9.3, color=MID_GRAY, bold=True)
        for source_index, (label, url) in enumerate(research["sources"]):
            if source_index:
                run = supplemental.add_run("；")
                _set_run_font(run, size=9.3, color=MID_GRAY)
            _add_hyperlink(supplemental, label, url, size=9.3)
        source = doc.add_paragraph()
        source.paragraph_format.space_after = Pt(8)
        run = source.add_run(f"来源：{case['source_name']}，发布日期{case['published_date']}。")
        _set_run_font(run, size=9.3, color=MID_GRAY)
        _add_hyperlink(source, "查看原文", case["url"], size=9.3)

    doc.add_heading("四、跨案风险特征", level=1)
    cross_risks = [
        ("涉华关联呈复合化", "中国来源货物、中国籍涉案人员、香港中转公司、日本转运路径和中国境内供应商等线索交织，应实行多字段联合识别。"),
        ("申报伪装与物理夹藏交织", "多品名误报、原产地换标、冷库板夹层、信号转换器夹金和机械设备藏毒同时出现，风险规则应覆盖申报文本、价格重量、设备结构和X光图像。"),
        ("三类运输渠道风险并存", "海运货柜单案规模大，航空旅客身份和航班信息较完整，邮包快件则高频、小批、主体分散，需要按渠道建立不同的串并模型。"),
        ("公开源发现与内部核查必须衔接", "多数新闻稿不披露集装箱号、提单号和完整企业名称，公开源情报适合作为风险发现入口，不能替代海关内部数据和执法协作渠道的二次核验。"),
    ]
    for title, body in cross_risks:
        _add_label_paragraph(doc, title + "：", body, after=8, size=11)

    doc.add_heading("五、下一步处置建议", level=1)
    recommendations = [
        ("建立强涉华案例线索台账", "将本报告已披露的人名、企业、网站、品牌、路线、货物和数量结构化录入，对别名、英文名和公司简称进行统一规范，逐项标记来源和可信度。"),
        ("开展单证反查", "对菲律宾电子烟和冷冻食品货柜、泰国换标货物、斯里兰卡冷库板夹烟等案件，获取箱号、提单号、收发货人和报关行信息，再与国内出口报关、舱单和物流数据比对。"),
        ("强化人员和企业关系串并", "优先核查Xin Wang、Gao Yong、Jing Tang Li、Kin Keung Ho、Lihua Owen Ma、DCP/PR、SJ、MT、TW以及Macropac、Megaspeed Services、Seg Metallic、PT TSL、PT TSI等主体，关联地址、电话、邮箱、账户和共同交易方。"),
        ("设置渠道化风险规则", "海运侧突出多柜集中到港、自由区换标和设备异常增重；航空侧突出多段中转、行李重量异常和贵金属未申报；邮包侧突出主纸箱拆分、预贴境内面单、科研用途标签和高频小包。"),
        ("建立处置反馈闭环", "对反查命中的企业和人员形成核查任务，记录是否发现同类申报、是否进入风险参数、是否开展查验及处置结果，并用于调整信息源权重、关键词和审核规则。"),
    ]
    for index, (title, body) in enumerate(recommendations, 1):
        _add_numbered_recommendation(doc, index, title, body)

    doc.add_page_break()
    doc.add_heading("附录：近30日境外进出口执法案例分类汇编", level=1)
    _add_body(
        doc,
        "共59条，按案件类型分组、组内按原文发布日期倒序排列。摘要优先采用系统中文译文，原始网页链接附后备查。",
        size=10,
        color=MID_GRAY,
    )
    grouped = defaultdict(list)
    for record in data["records"]:
        grouped[record["category"]].append(record)
    overall_index = 0
    for category in data["category_order"]:
        records = grouped.get(category, [])
        if not records:
            continue
        doc.add_heading(f"{category}（{len(records)}条）", level=2)
        for record in records:
            overall_index += 1
            heading = doc.add_heading(f"{overall_index}. {record['title']}", level=3)
            heading.paragraph_format.keep_with_next = True
            _add_body(doc, record["summary"], after=4, size=10.2)
            meta = doc.add_paragraph()
            meta.paragraph_format.space_after = Pt(2)
            run = meta.add_run(
                f"发布日期：{record['published_date']}  |  国家或地区：{record['jurisdiction']}  |  "
                f"涉华等级：{record['level_label']}  |  执法机关：{record['authority']}"
            )
            _set_run_font(run, size=9, color=MID_GRAY)
            source = doc.add_paragraph()
            source.paragraph_format.space_after = Pt(7)
            run = source.add_run("来源：")
            _set_run_font(run, size=9, color=MID_GRAY)
            _add_hyperlink(source, record["source_name"], record["url"], size=9)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(str(output_path))


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_enforcement_report_docx.py INPUT_JSON OUTPUT_DOCX")
    data = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    build(data, Path(sys.argv[2]))


if __name__ == "__main__":
    main()
