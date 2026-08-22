"""Import operator-provided source and article workbooks as database user data."""
from __future__ import annotations

import fnmatch
import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlsplit, urlunsplit

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import SourceChannel, SourceConfig, Topic
from app.source_inventory_import import normalize_text, normalize_url


_SITE_HEADERS = frozenset({"机构名称", "网站", "国家/地区"})
_TITLE_HEADERS = ("标题", "文章标题")
_SUMMARY_HEADERS = ("主要内容摘要", "主要内容")
_URL_HEADERS = ("来源链接", "链接", "原文链接")
_DATE_HEADERS = ("发布日期", "时间", "日期")
_PUBLISHER_HEADERS = ("来源网站名称", "来源网站", "来源")


@dataclass(frozen=True, slots=True)
class TopicSpec:
    id: str
    name: str
    file_patterns: tuple[str, ...]
    sheet_patterns: tuple[str, ...]
    values: dict[str, Any] = field(compare=False)


@dataclass(frozen=True, slots=True)
class DirectorySite:
    name: str
    url: str
    region: str
    country: str
    access_observation: str
    robots_observation: str
    source_file: str


@dataclass(frozen=True, slots=True)
class ArticleTarget:
    topic_id: str
    title: str
    summary: str
    publisher: str
    url: str
    published_at: str | None
    category: str | None
    source_files: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class WorkbookDataset:
    topics: tuple[TopicSpec, ...]
    sites: tuple[DirectorySite, ...]
    targets: tuple[ArticleTarget, ...]


@dataclass(frozen=True, slots=True)
class ImportSummary:
    workbooks: int
    sites_parsed: int
    targets_parsed: int
    sources_created: int
    sources_reused: int
    topics_created: int
    topics_updated: int
    targets_added: int


def parse_workbook_directory(root: Path, manifest_path: Path) -> WorkbookDataset:
    """Parse every XLSX using manifest-provided topic routing, then deduplicate."""
    topics = _load_topic_specs(manifest_path)
    site_map: dict[str, DirectorySite] = {}
    target_map: dict[tuple[str, str], ArticleTarget] = {}
    for path in sorted(root.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        workbook = load_workbook(path, data_only=True, read_only=False)
        for sheet in workbook.worksheets:
            header = _find_header(sheet)
            if header is None:
                continue
            row_number, columns = header
            if _SITE_HEADERS.issubset(columns):
                for site in _parse_site_rows(path, sheet, row_number, columns):
                    site_map.setdefault(normalize_url(site.url), site)
                continue
            topic = _match_topic(topics, path.name, sheet.title)
            if topic is None or not _is_article_header(columns):
                continue
            for target in _parse_article_rows(path, sheet, row_number, columns, topic.id):
                key = (topic.id, normalize_url(target.url))
                previous = target_map.get(key)
                target_map[key] = target if previous is None else ArticleTarget(
                    topic_id=previous.topic_id,
                    title=previous.title or target.title,
                    summary=previous.summary or target.summary,
                    publisher=previous.publisher or target.publisher,
                    url=previous.url,
                    published_at=previous.published_at or target.published_at,
                    category=previous.category or target.category,
                    source_files=tuple(dict.fromkeys([
                        *previous.source_files, *target.source_files,
                    ])),
                )
    return WorkbookDataset(
        topics=topics,
        sites=tuple(site_map[key] for key in sorted(site_map)),
        targets=tuple(target_map[key] for key in sorted(target_map)),
    )


def import_workbook_directory(
    db: Session, root: Path, manifest_path: Path,
) -> ImportSummary:
    """Idempotently merge parsed user data without changing operator compliance."""
    dataset = parse_workbook_directory(root, manifest_path)
    sources = list(db.query(SourceConfig).all())
    site_source_ids: dict[str, str] = {}
    created = reused = 0
    for site in dataset.sites:
        source = _find_site_source(sources, site)
        if source is None:
            source = _new_site_source(site)
            db.add(source)
            sources = [*sources, source]
            created += 1
        else:
            _record_import_provenance(source, site)
            reused += 1
        site_source_ids[normalize_url(site.url)] = source.id
    db.flush()

    target_sources: dict[tuple[str, str], str] = {}
    for target in dataset.targets:
        source = _find_target_source(sources, target)
        if source is None:
            source = _new_target_source(target)
            db.add(source)
            sources = [*sources, source]
            created += 1
        target_sources[(target.topic_id, normalize_url(target.url))] = source.id
    db.flush()

    topics_created = topics_updated = targets_added = 0
    for spec in dataset.topics:
        matching_targets = [target for target in dataset.targets if target.topic_id == spec.id]
        topic = db.get(Topic, spec.id)
        if topic is None:
            topic = Topic(id=spec.id, name=spec.name)
            db.add(topic)
            topics_created += 1
        else:
            topics_updated += 1
        previous_targets = list(topic.target_urls or [])
        topic.target_urls = _merge_target_urls(
            previous_targets, [target.url for target in matching_targets],
        )
        targets_added += len(topic.target_urls) - len(previous_targets)
        topic.source_ids = _topic_source_ids(
            topic, spec, dataset.sites, site_source_ids,
            matching_targets, target_sources,
        )
        _apply_topic_values(topic, spec.values)
    db.flush()
    return ImportSummary(
        workbooks=len(list(root.glob("*.xlsx"))),
        sites_parsed=len(dataset.sites),
        targets_parsed=len(dataset.targets),
        sources_created=created,
        sources_reused=reused,
        topics_created=topics_created,
        topics_updated=topics_updated,
        targets_added=targets_added,
    )


def _load_topic_specs(path: Path) -> tuple[TopicSpec, ...]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    specs = []
    for row in payload.get("topics", []):
        specs.append(TopicSpec(
            id=str(row["id"]), name=str(row["name"]),
            file_patterns=tuple(row.get("file_patterns") or ("*.xlsx",)),
            sheet_patterns=tuple(row.get("sheet_patterns") or ("*",)),
            values=dict(row),
        ))
    return tuple(specs)


def _find_header(sheet) -> tuple[int, dict[str, int]] | None:
    for row in sheet.iter_rows(min_row=1, max_row=min(12, sheet.max_row)):
        columns = {
            str(cell.value or "").strip(): index
            for index, cell in enumerate(row)
            if str(cell.value or "").strip()
        }
        if _SITE_HEADERS.issubset(columns) or _is_article_header(columns):
            return row[0].row, columns
    return None


def _is_article_header(columns: dict[str, int]) -> bool:
    return bool(
        _first_column(columns, _TITLE_HEADERS) is not None
        and _first_column(columns, _URL_HEADERS) is not None
    )


def _parse_site_rows(path, sheet, header_row, columns):
    for row in sheet.iter_rows(min_row=header_row + 1):
        name = _cell_text(row[columns["机构名称"]])
        url = _cell_url(row[columns["网站"]])
        if name and url:
            yield DirectorySite(
                name=name, url=url, region=sheet.title,
                country=_cell_text(row[columns["国家/地区"]]),
                access_observation=_cell_text(row[columns.get("中国大陆访问", 0)]),
                robots_observation=_cell_text(row[columns.get("反爬虫机制", 0)]),
                source_file=path.name,
            )


def _parse_article_rows(path, sheet, header_row, columns, topic_id):
    title_col = _first_column(columns, _TITLE_HEADERS)
    url_col = _first_column(columns, _URL_HEADERS)
    for row in sheet.iter_rows(min_row=header_row + 1):
        title = _cell_text(row[title_col])
        url = _cell_url(row[url_col])
        if not title or not url:
            continue
        yield ArticleTarget(
            topic_id=topic_id, title=title, url=url,
            summary=_optional_cell(row, columns, _SUMMARY_HEADERS),
            publisher=_optional_cell(row, columns, _PUBLISHER_HEADERS),
            published_at=_date_cell(row, columns),
            category=_optional_cell(row, columns, ("类别",)) or None,
            source_files=(path.name,),
        )


def _match_topic(topics, filename, sheet_name):
    return next((spec for spec in topics if (
        any(fnmatch.fnmatch(filename, pattern) for pattern in spec.file_patterns)
        and any(fnmatch.fnmatch(sheet_name, pattern) for pattern in spec.sheet_patterns)
    )), None)


def _find_site_source(sources, site):
    key = normalize_url(site.url)
    matches = [source for source in sources if key in _source_urls(source)]
    return _best_source(matches, site.name)


def _find_target_source(sources, target):
    origin = _origin_url(target.url).casefold()
    matches = [source for source in sources if (
        _channel(source) == "web_scrape"
        and origin in {
            _origin_url(url).casefold()
            for url in (source.base_url, source.homepage_url) if url
        }
    )]
    return _best_source(matches, target.publisher)


def _best_source(sources, name):
    if not sources:
        return None
    normalized = normalize_text(name)
    return max(sources, key=lambda source: (
        normalize_text(source.name) == normalized,
        normalized in normalize_text(source.name) or normalize_text(source.name) in normalized,
        str(source.verification_status or "").startswith("verified"),
        _channel(source) == "web_scrape",
        -len(source.base_url or ""),
        source.id,
    ))


def _new_site_source(site: DirectorySite) -> SourceConfig:
    profile = _site_profile(site)
    return SourceConfig(
        id=_stable_id("user-site", normalize_url(site.url)), name=site.name[:200],
        description=f"用户工作簿导入：{site.region} / {site.country}",
        channel=SourceChannel.WEB_SCRAPE, source_group=_site_group(site.region),
        is_active=site.access_observation == "可访问", is_configured=True,
        base_url=site.url, homepage_url=_origin_url(site.url),
        rate_limit_rps=0.2, crawl_delay_seconds=_crawl_delay(site.robots_observation),
        max_items_per_run=20, country_focus=[site.country] if site.country else None,
        default_categories=[site.region], verification_status="imported_unverified",
        robots_status=_robots_status(site.robots_observation), terms_status="unverified",
        llm_ingest_allowed=False, collection_profile=profile,
        compliance_note="用户目录数据；须在系统内完成人工合规审核后方可自动采集。",
    )


def _new_target_source(target: ArticleTarget) -> SourceConfig:
    origin = _origin_url(target.url)
    return SourceConfig(
        id=_stable_id("user-target", origin),
        name=(target.publisher or _host(target.url) or "用户资讯目标")[:200],
        description="由用户工作簿中的原文链接创建，用于读取指定详情页正文。",
        channel=SourceChannel.WEB_SCRAPE, source_group="government_igo",
        is_active=True, is_configured=True, base_url=origin, homepage_url=origin,
        auth_config={"target_only": True}, rate_limit_rps=0.2,
        crawl_delay_seconds=8, max_items_per_run=20,
        verification_status="imported_unverified", robots_status="unverified",
        terms_status="unverified", llm_ingest_allowed=False,
        collection_profile={
            "origin": "user_workbook", "source_files": list(target.source_files),
            "target_count": 1,
        },
        compliance_note="用户提供原文目标；须完成来源、robots、条款与LLM使用许可审核。",
    )


def _record_import_provenance(source, site):
    profile = dict(source.collection_profile or {})
    imports = list(profile.get("user_workbook_imports") or [])
    record = _site_profile(site)
    source.collection_profile = {
        **profile,
        "user_workbook_imports": [*imports, record] if record not in imports else imports,
    }


def _site_profile(site):
    return {
        "origin": "user_workbook", "source_file": site.source_file,
        "region": site.region, "country": site.country,
        "access_observation": site.access_observation,
        "robots_observation": site.robots_observation,
    }


def _topic_source_ids(topic, spec, sites, site_source_ids, targets, target_sources):
    ids = list(topic.source_ids or [])
    ids.extend(target_sources[(target.topic_id, normalize_url(target.url))] for target in targets)
    filters = set(spec.values.get("source_country_filters") or [])
    for site in sites:
        if spec.values.get("bind_all_directory_sources") or site.country in filters:
            ids.append(site_source_ids[normalize_url(site.url)])
    return list(dict.fromkeys(ids))


def _apply_topic_values(topic, values):
    list_fields = (
        "keywords", "synonyms", "exclude_keywords", "categories",
        "focus_countries", "focus_languages",
    )
    for name in list_fields:
        if name in values:
            current = list(getattr(topic, name, None) or [])
            topic_value = list(values[name] or [])
            setattr(topic, name, list(dict.fromkeys([*current, *topic_value])))
    scalar_fields = (
        "description", "description_prompt", "collect_window_days",
        "collection_policy", "is_active",
    )
    for name in scalar_fields:
        if name in values and getattr(topic, name, None) in (None, [], ""):
            setattr(topic, name, values[name])
    if "target_urls_mode" in values:
        topic.target_urls_mode = values["target_urls_mode"]
    if topic.keywords is None:
        topic.keywords = [topic.name]


def _source_urls(source):
    base = str(source.base_url or source.homepage_url or "")
    endpoint = str(source.api_endpoint or "")
    resolved_endpoint = (
        endpoint if endpoint.startswith(("http://", "https://"))
        else urljoin(f"{base.rstrip('/')}/", endpoint.lstrip("/")) if endpoint and base
        else ""
    )
    values = [
        source.base_url, source.homepage_url, resolved_endpoint,
        *(source.discovery_urls or []),
    ]
    return {normalize_url(url) for url in values if url}


def _merge_target_urls(current, incoming):
    merged: dict[str, str] = {}
    for url in [*current, *incoming]:
        key = normalize_url(url)
        if key:
            merged.setdefault(key, url)
    return list(merged.values())


def _channel(source):
    return str(getattr(source.channel, "value", source.channel) or "").casefold()


def _origin_url(value):
    parsed = urlsplit(value)
    return urlunsplit((parsed.scheme, parsed.netloc, "/", "", ""))


def _host(value):
    return (urlsplit(value).hostname or "").casefold().removeprefix("www.")


def _stable_id(prefix, value):
    return f"{prefix}-{hashlib.sha256(value.encode()).hexdigest()[:12]}"


def _site_group(region):
    return "news_risk_media" if "智库" in region else "government_igo"


def _robots_status(value):
    if re.search(r"全站禁爬|封锁|搜索禁爬", value):
        return "blocked_import_observation"
    return "observed_restrictions" if value.startswith("有") else "unverified"


def _crawl_delay(value):
    match = re.search(r"Crawl-delay\s*[:：]\s*(\d+)", value, re.I)
    return int(match.group(1)) if match else 8


def _first_column(columns, names):
    return next((columns[name] for name in names if name in columns), None)


def _optional_cell(row, columns, names):
    index = _first_column(columns, names)
    return _cell_text(row[index]) if index is not None else ""


def _date_cell(row, columns):
    index = _first_column(columns, _DATE_HEADERS)
    if index is None:
        return None
    value = row[index].value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return _cell_text(row[index]) or None


def _cell_text(cell):
    return str(cell.value or "").strip()


def _cell_url(cell):
    target = cell.hyperlink.target if cell.hyperlink else None
    value = target if target and str(target).startswith("http") else cell.value
    text = str(value or "").strip()
    return text if text.startswith(("http://", "https://")) else ""


__all__ = [
    "ArticleTarget", "DirectorySite", "ImportSummary", "WorkbookDataset",
    "import_workbook_directory", "parse_workbook_directory",
]
