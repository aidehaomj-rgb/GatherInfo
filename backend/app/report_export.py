"""
Report Export — render a generated Report's markdown into MD / HTML / DOCX / PDF
files on disk, organized into a date-based subfolder.

Design notes:
    - CJK support: PDF uses reportlab's built-in CID font "STSong-Light" (no external
      font files needed). DOCX uses python-docx defaults which already handle CJK.
    - Heavy deps (reportlab, python-docx) are imported lazily so a missing optional
      dependency only disables that one format instead of breaking the whole app.
    - Output layout: {root}/{date_pattern}/{safe_title}.{ext}
      root default = <repo>/data/reports
"""
from __future__ import annotations

import html as _html
import logging
import os
import re
import tempfile
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from app.database import DATA_DIR
from app.models import Report, SystemConfig, Topic

SUPPORTED_FORMATS = ("md", "html", "docx", "pdf")
logger = logging.getLogger(__name__)
DEFAULT_REPORT_ROOT = os.path.realpath(os.path.join(DATA_DIR, "reports"))


# ── Public API ──────────────────────────────────────────────────────────────

def export_report(
    report: Report,
    system: SystemConfig | None,
    topic: Topic | None,
    formats: list[str] | None = None,
) -> dict[str, str]:
    """Render `report` to all configured formats. Returns {format: abs_path}.

    Existing, valid files are retained. Failed formats are logged and omitted.
    """
    target_formats = _resolve_formats(system) if formats is None else _valid_formats(formats)
    title = _resolve_title(report, system, topic)
    safe_title = _safe_filename(title)

    root = resolve_report_output_dir(
        system.report_output_dir if system and system.report_output_dir else None,
    )
    pattern = (system.report_dir_pattern if system and system.report_dir_pattern else "%Y-%m-%d")
    subdir = _safe_report_subdir(pattern)
    out_dir = os.path.join(os.path.abspath(root), subdir)
    os.makedirs(out_dir, exist_ok=True)

    body = report.content or ""
    out_files = valid_output_files(report.output_files)

    for fmt in target_formats:
        path = os.path.join(out_dir, f"{safe_title}.{fmt}")
        fd, temporary_path = tempfile.mkstemp(prefix=".report-", suffix=".tmp", dir=out_dir)
        os.close(fd)
        try:
            if fmt == "md":
                _write_md(temporary_path, title, body)
            elif fmt == "html":
                _write_html(temporary_path, title, body)
            elif fmt == "docx":
                _write_docx(temporary_path, title, body)
            elif fmt == "pdf":
                _write_pdf(temporary_path, title, body)
            else:
                continue
            os.replace(temporary_path, path)
            out_files[fmt] = path
        except ImportError:
            logger.warning("Report %s export skipped for %s: optional dependency unavailable", report.id, fmt)
            continue
        except Exception as exc:
            logger.exception("Report %s export failed for %s: %s", report.id, fmt, exc)
            continue
        finally:
            if os.path.exists(temporary_path):
                os.unlink(temporary_path)

    report.output_files = out_files
    report.output_dir = out_dir
    return out_files


# ── Helpers ─────────────────────────────────────────────────────────────────

def _resolve_formats(system: SystemConfig | None) -> list[str]:
    fmts = (system.report_formats if system and system.report_formats else None) or list(SUPPORTED_FORMATS)
    return _valid_formats(fmts)


def _valid_formats(formats: list[str]) -> list[str]:
    return list(dict.fromkeys(str(fmt).lower() for fmt in formats if str(fmt).lower() in SUPPORTED_FORMATS))


def normalize_report_output_dir(value: str | None) -> str | None:
    """Normalize user-entered output paths and tolerate copied shell quotes."""
    if not value:
        return None
    normalized = value.strip()
    if len(normalized) >= 2 and normalized[0] == normalized[-1] and normalized[0] in {"'", '"'}:
        normalized = normalized[1:-1].strip()
    return os.path.abspath(os.path.expanduser(normalized)) if normalized else None


def approved_report_roots() -> tuple[str, ...]:
    configured = tuple(
        os.path.realpath(os.path.abspath(os.path.expanduser(value.strip())))
        for value in os.getenv("REPORT_OUTPUT_ROOTS", "").split(os.pathsep)
        if value.strip()
    )
    return tuple(dict.fromkeys((DEFAULT_REPORT_ROOT, *configured)))


def is_approved_report_path(path: str | None, *, require_file: bool = False) -> bool:
    if not path or not os.path.isabs(path) or os.path.islink(path):
        return False
    resolved = Path(os.path.realpath(path))
    allowed = any(
        resolved == Path(root) or resolved.is_relative_to(Path(root))
        for root in approved_report_roots()
    )
    return allowed and (not require_file or resolved.is_file())


def resolve_report_output_dir(value: str | None) -> str:
    candidate = normalize_report_output_dir(value) or DEFAULT_REPORT_ROOT
    if not is_approved_report_path(candidate):
        raise ValueError(
            "报告输出目录不在服务器批准范围；请通过 REPORT_OUTPUT_ROOTS 显式批准。"
        )
    return candidate


def valid_output_files(files: dict | None) -> dict[str, str]:
    """Keep only download-ready absolute paths from persisted export metadata."""
    if not isinstance(files, dict):
        return {}
    return {
        fmt: path for fmt, path in files.items()
        if fmt in SUPPORTED_FORMATS and isinstance(path, str)
        and is_approved_report_path(path, require_file=True)
    }


def _resolve_title(report: Report, system: SystemConfig | None, topic: Topic | None) -> str:
    fmt = (system.report_title_format if system and system.report_title_format else None) \
        or "{title}_{date}"
    topic_name = (topic.name if topic else None) or report.topic_id or "报告"
    date_str = datetime.now().strftime("%Y-%m-%d")
    try:
        return fmt.format(topic=topic_name, date=date_str, title=report.title or topic_name)
    except (KeyError, IndexError, ValueError):
        return report.title or f"{topic_name}_{date_str}"


_UNSAFE = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def _safe_filename(name: str) -> str:
    cleaned = _UNSAFE.sub("_", name).strip(" .")
    cleaned = cleaned[:120] if len(cleaned) > 120 else cleaned
    return cleaned or "report"


def _safe_report_subdir(pattern: str) -> str:
    """Preserve date nesting while preventing traversal outside the report root."""
    rendered = datetime.now().strftime(pattern or "%Y-%m-%d")
    parts = [
        _safe_filename(part)
        for part in re.split(r"[\\/]+", rendered)
        if part.strip() not in {"", ".", ".."}
    ]
    return os.path.join(*parts) if parts else datetime.now().strftime("%Y-%m-%d")


def _write_md(path: str, title: str, body: str) -> None:
    text = body if body.lstrip().startswith("#") else f"# {title}\n\n{body}"
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)


def _write_html(path: str, title: str, body: str) -> None:
    content_html = _markdown_to_html(body)
    doc = (
        "<!DOCTYPE html>\n<html lang=\"zh-CN\">\n<head>\n<meta charset=\"utf-8\">\n"
        "<meta http-equiv=\"Content-Security-Policy\" content=\"default-src 'none'; style-src 'unsafe-inline'; img-src https: data:;\">\n"
        f"<title>{_html.escape(title)}</title>\n"
        "<style>\n"
        "body{font-family:-apple-system,'Segoe UI','PingFang SC','Microsoft YaHei',sans-serif;"
        "max-width:860px;margin:40px auto;padding:0 20px;line-height:1.7;color:#1f2937;}\n"
        "h1,h2,h3{color:#111827;margin-top:1.4em;}\n"
        "code{background:#f3f4f6;padding:2px 5px;border-radius:4px;}\n"
        "pre{background:#f3f4f6;padding:12px;border-radius:8px;overflow:auto;}\n"
        "blockquote{border-left:3px solid #d1d5db;margin:0;padding-left:14px;color:#4b5563;}\n"
        "table{border-collapse:collapse;}td,th{border:1px solid #d1d5db;padding:6px 10px;}\n"
        "</style>\n</head>\n<body>\n"
        f"{content_html}\n</body>\n</html>\n"
    )
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(doc)


def _write_docx(path: str, title: str, body: str) -> None:
    from docx import Document  # lazy import

    doc = Document()
    doc.add_heading(title, level=0)
    for line in body.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\.\s", stripped):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", stripped), style="List Number")
        else:
            doc.add_paragraph(_strip_inline_md(stripped))
    doc.save(path)


def _write_pdf(path: str, title: str, body: str) -> None:
    from reportlab.lib.pagesizes import A4  # lazy import
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    font_name = "STSong-Light"
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    except Exception:
        font_name = "Helvetica"

    base = getSampleStyleSheet()
    body_style = ParagraphStyle("Body", parent=base["BodyText"], fontName=font_name,
                                fontSize=10.5, leading=16)
    h1 = ParagraphStyle("H1", parent=base["Heading1"], fontName=font_name, fontSize=18, leading=24)
    h2 = ParagraphStyle("H2", parent=base["Heading2"], fontName=font_name, fontSize=14, leading=20)
    h3 = ParagraphStyle("H3", parent=base["Heading3"], fontName=font_name, fontSize=12, leading=18)
    title_style = ParagraphStyle("Title", parent=base["Title"], fontName=font_name, fontSize=22, leading=28)

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    flow = [Paragraph(_html.escape(title), title_style), Spacer(1, 8)]
    for line in body.splitlines():
        stripped = line.strip()
        if not stripped:
            flow.append(Spacer(1, 6))
            continue
        if stripped.startswith("### "):
            flow.append(Paragraph(_html.escape(stripped[4:]), h3))
        elif stripped.startswith("## "):
            flow.append(Paragraph(_html.escape(stripped[3:]), h2))
        elif stripped.startswith("# "):
            flow.append(Paragraph(_html.escape(stripped[2:]), h1))
        elif stripped.startswith(("- ", "* ")):
            flow.append(Paragraph("• " + _html.escape(_strip_inline_md(stripped[2:])), body_style))
        else:
            flow.append(Paragraph(_html.escape(_strip_inline_md(stripped)), body_style))
    doc.build(flow)


# ── Minimal markdown helpers ─────────────────────────────────────────────────

def _strip_inline_md(text: str) -> str:
    """Remove the most common inline markdown markers for plain-text renderers."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", text)
    return text


def _markdown_to_html(body: str) -> str:
    """Very small markdown→HTML for headings, lists, bold/italic/code and paragraphs."""
    lines = body.splitlines()
    out: list[str] = []
    in_list = False

    def close_list() -> None:
        nonlocal in_list
        if in_list:
            out.append("</ul>")
            in_list = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            close_list()
            continue
        if stripped.startswith("### "):
            close_list(); out.append(f"<h3>{_inline_html(stripped[4:])}</h3>")
        elif stripped.startswith("## "):
            close_list(); out.append(f"<h2>{_inline_html(stripped[3:])}</h2>")
        elif stripped.startswith("# "):
            close_list(); out.append(f"<h1>{_inline_html(stripped[2:])}</h1>")
        elif stripped.startswith(("- ", "* ")):
            if not in_list:
                out.append("<ul>"); in_list = True
            out.append(f"<li>{_inline_html(stripped[2:])}</li>")
        else:
            close_list(); out.append(f"<p>{_inline_html(stripped)}</p>")
    close_list()
    return "\n".join(out)


def _inline_html(text: str) -> str:
    text = _html.escape(text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"\*(.+?)\*", r"<em>\1</em>", text)
    text = re.sub(r"`(.+?)`", r"<code>\1</code>", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", _safe_html_link, text)
    return text


def _safe_html_link(match: re.Match[str]) -> str:
    label, escaped_url = match.group(1), match.group(2)
    scheme = urlsplit(_html.unescape(escaped_url)).scheme.casefold()
    if scheme not in {"http", "https"}:
        return label
    return (
        f'<a href="{escaped_url}" rel="noopener noreferrer" '
        f'target="_blank">{label}</a>'
    )
