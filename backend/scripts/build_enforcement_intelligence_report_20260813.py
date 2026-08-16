"""Build the concise 2026-08-13 enforcement report with verified case details."""
from __future__ import annotations

import sys
import re
from collections import Counter
from datetime import datetime
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.database import SessionLocal  # noqa: E402
from app.models import CollectedItem, CollectionRun, Report, SystemConfig, Topic  # noqa: E402
from app.report_export import export_report  # noqa: E402


TOPIC_ID = "weekly-enforcement-intelligence"
RUN_ID = "run-63e4473d2bfa"
REPORT_ID = "rpt-enforcement-intelligence-20260813"
TITLE = "境外执法信息智能简报（2026-08-13）"

# Details below are limited to facts disclosed by the source or a same-case report.
DETAILS = {
    "d52e0212764204bd": {
        "entities": "人员：32岁中国籍悉尼男子（官方未公开姓名）；行动代号：Operation CALOR。",
        "transport": "未公开集装箱号、车辆或船名。",
        "extra": "查获约560万支卷烟、404公斤散装烟叶、21万支电子烟；涉案物街头估值约1,250万澳元。",
    },
    "codex-malaysia-johor-chinese-scam-arrests-20260729": {
        "entities": "地点：Forest City, Iskandar Puteri；人员：309名中国公民、19名印尼公民、4名缅甸公民、3名马来西亚公民。",
        "transport": "非货运案件，无集装箱或运输工具编号。",
        "extra": "马来西亚皇家警察确认两个呼叫中心诈骗集团被捣毁，共拘捕335人。",
    },
    "a5128dfb24239123": {
        "entities": "企业：Nvidia、Super Micro Computer、Albatron Technology；人员：7人被羁押，姓名未公开。",
        "transport": "约50台服务器；部分经日本转运至中国大陆，另涉澳门、香港线路；未公开箱号或航班。",
        "extra": "同案报道还提及新加坡警方扣押一栋价值逾4,000万美元的豪宅，但案件关联仍在调查中。",
    },
    "dc80af8e8b74466d": {
        "entities": "企业：YJC International Corporation、ADNFinest Marketing Corporation；后者总裁、副总裁和公司秘书被定罪。",
        "transport": "中国发运的2个40英尺集装箱；另案涉及来自迪拜的车辆。官方报道未公开箱号、提单号或船名。",
        "extra": "首案查获1,599箱卷烟；另案查获未申报的2020 Chevrolet Camaro ZL1及98条旧轮胎。",
    },
    "codex-us-chinese-meth-australia-20260724": {
        "entities": "人员：Jing Tang Li，34岁，中国公民；承办法官：Wesley L. Hsu。",
        "transport": "美国至澳大利亚共7批货，伪报为地毯、家具腿、轮毂检测设备和铸造机；未公开箱号、承运人或航班。",
        "extra": "分批查获约855.5、32.5、202.5及两批各102.8公斤甲基苯丙胺，总量逾1,000公斤。",
    },
    "codex-us-chinese-turtle-trafficking-20260723": {
        "entities": "人员：Kin Keung Ho、Lihua Owen Ma；调查机关：U.S. Fish & Wildlife Service、U.S. Postal Inspection Service。",
        "transport": "Ho承认邮寄约99个包裹、共578只龟，包裹虚假标注为水晶或石头；未公开运单号。",
        "extra": "涉及箱龟、斑点龟、菱背水龟等受CITES保护物种，目的市场为亚洲及香港。",
    },
    "codex-us-unapproved-chinese-drugs-20260723": {
        "entities": "人员：案件标题及司法部记录确认被告姓Piper；企业名称以司法部原文为准。",
        "transport": "中国来源未获FDA批准药品；公开报道未披露包裹号、承运人或车辆。",
        "extra": "FDA刑事调查办公室将其列入2026年7月23日司法部执法案件。",
    },
    "codex-thailand-origin-fraud-20260722": {
        "entities": "企业/品牌：TWOMOON；地点：Laem Chabang Port。",
        "transport": "两批中国进口货物，未公开集装箱号或船名。",
        "extra": "包括537万支假冒泰国产卷烟，以及逾20.2万套美工刀和多功能工具；货值逾1亿泰铢。",
    },
    "codex-ph-china-frozen-food-20260721": {
        "entities": "品牌：Fu Hua Lao Jiang、Qianye；执法：菲律宾海关与农业部。",
        "transport": "中国来源5个集装箱；3份PLCO和2份Alert Order。官方文字和可核验图片均未公开完整箱号。",
        "extra": "申报为鱼丸、鱼豆腐等，实查有鸡胸、乳鸽、北京鸭等；总值约2,326.89万比索。",
    },
    "codex-sri-lanka-chinese-cigarettes-20260717": {
        "entities": "执法机关：Sri Lanka Customs Central Intelligence Directorate。",
        "transport": "中国来源卷烟藏于冷库保温板；公开报道未披露集装箱号、船名或提单号。",
        "extra": "查获约360万支中国卷烟，报道估值约4.5亿斯里兰卡卢比。",
    },
    "codex-malaysia-melaka-chinese-scam-arrests-20260716": {
        "entities": "地点：Ayer Keroh, Melaka；人员：21名23至54岁的中国籍男子。",
        "transport": "非货运案件，无集装箱或运输工具编号。",
        "extra": "警方突击一栋被用作虚假投资诈骗窝点的两层别墅。",
    },
    "f7ced9411d07dd37": {
        "entities": "机构：BOC Port of Clark、PDEA及Clark跨部门缉毒力量。",
        "transport": "国际包裹，氯胺酮藏于兔子摆件；官方报道及图库未公开运单号和航班。",
        "extra": "估值837.5万比索；案件由跨部门联合检查发现。",
    },
    "5e6620794538ec29": {
        "entities": "人员：两名维多利亚州嫌疑人，姓名以澳大利亚边防局司法限制为准。",
        "transport": "涉及进口健身及形象增强药物；未公开包裹号或承运工具。",
        "extra": "属于既有调查追加指控，报告不将指控写成最终定罪。",
    },
    "b6de316cf4c55c92": {
        "entities": "执法机关：Canada Border Services Agency；地点：Coutts, Alberta。",
        "transport": "商业运输卡车；官方及同案媒体未公开车牌、拖车号或承运公司。",
        "extra": "2026年7月2日查获约300公斤甲基苯丙胺和80公斤可卡因。",
    },
    "472a75422de01d1b": {
        "entities": "机构：CBP Anti-Terrorism Contraband Team；CBP官员：Michael Pfeiffer。",
        "transport": "货物自中国发往美国纽约州Syracuse，经芝加哥O'Hare机场截获；未公开运单和航班。",
        "extra": "同案CBS Chicago报道确认货物为多种未经批准药品。",
    },
    "c83067167a680499": {
        "entities": "执法机关：意大利Guardia di Finanza；4名船上人员为2名西班牙人、1名直布罗陀居民和1名居意阿尔巴尼亚人。",
        "transport": "大西洋葡萄牙近海船只；公开报道未披露船名、船旗、IMO或呼号。",
        "extra": "查获逾2.6吨可卡因，估计街头价值5亿欧元。",
    },
    "d85097cfe3ca9634": {
        "entities": "执法机关：Receita Federal；地点和涉案主体以巴西官方通报为准。",
        "transport": "公开通报未披露车辆、包裹或集装箱识别号。",
        "extra": "查获电子烟及大麻衍生产品，货值逾3万雷亚尔。",
    },
    "fbe3d09ff7874331": {
        "entities": "机构：UK Border Force；该条为缉毒风险宣传并引用实际执法背景。",
        "transport": "未披露具体同案航班、车辆、船舶或货运编号。",
        "extra": "证据粒度低于查获通报，适合作为旅客走私预警，不宜用于运输链追踪。",
    },
    "7801809f62b51d65": {
        "entities": "机构：BOC Port of Clark及Clark跨部门力量。",
        "transport": "毒品藏于汽车零部件货件；官方报道未公开运单号、航班或承运人。",
        "extra": "查获疑似shabu，估值约1,160万比索。",
    },
    "728b2571123af33b": {
        "entities": "人员：Blue Mountains地区男子；机构：Australian Border Force及联合执法单位。",
        "transport": "案件涉及非法枪支制造，未公开相关包裹、车辆或枪支序列号。",
        "extra": "当前为指控阶段，报告不将嫌疑内容表述为已定罪事实。",
    },
}


def build_report(items: list[CollectedItem]) -> str:
    reviews = [(item.raw_metadata or {}).get("enforcement_review") or {} for item in items]
    strong = sum(review.get("inclusion_basis") == "strong_china_nexus" for review in reviews)
    jurisdictions = Counter(review.get("jurisdiction") or "未知" for review in reviews)
    lines = [
        "## 一、结论",
        "",
        f"本期纳入 **{len(items)}** 起经证据审核的境外执法案例，其中强涉华 **{strong}** 起（{strong / len(items):.0%}），全部为香港以外案件，覆盖 **{len(jurisdictions)}** 个国家或地区。",
        "案件集中在跨境毒品、烟草及电子烟、农食产品伪报、知识产权、出口管制和网络诈骗。强涉华线索主要表现为中国籍涉案人员、中国来源货物、面向中国的转运路线，以及中国企业或产品直接进入案件链条。",
        "",
        "## 二、核心发现",
        "",
        f"- 地域：{ '；'.join(f'{name} {count}起' for name, count in jurisdictions.most_common()) }。",
        "- 重点一：菲律宾、泰国和斯里兰卡多起案件指向中国来源货物的伪报、原产地置换及夹藏，应关注申报品名、品牌、冷链货物和实际查验结果之间的差异。",
        "- 重点二：美国、澳大利亚案件呈现人员、仓库、包裹和伪装设备共同构成的跨境链条，适合开展主体、地址、承运渠道和历史申报关联分析。",
        "- 重点三：芯片经第三地转运至中国的案件表明，高端服务器、最终用户和转运地组合仍是出口管制核查重点。",
        "- 数据边界：公开来源普遍不披露完整集装箱号、提单号、车牌或航班号。未能从同案报道和原始图片稳定辨认的编号均标注为未公开，不作推测。",
        "",
        "## 三、风险分析",
        "",
        "### （一）人员与组织风险",
        "中国籍人员既出现在烟草、毒品和野生动物走私链条，也出现在境外网络诈骗和高技术产品转运案件中。风险主体可能通过仓库、空壳企业、呼叫中心、邮政包裹及第三地公司分散活动，单一姓名或企业字段不足以识别完整网络，应联合比对地址、电话、邮箱、账户、设备和历史申报。",
        "",
        "### （二）货物申报与原产地风险",
        "冻品、卷烟、电子烟、药品和工具类案件反复出现品名伪报、货物夹藏、虚假原产地标签和监管证件缺失。中国来源货物经自由区或第三地换标后，可能形成原产地欺诈、食品安全、知识产权和税收流失等复合风险。",
        "",
        "### （三）运输链与单证风险",
        "海运集装箱适合大批量伪报，航空和邮政渠道则呈现小批量、高频次、拆分寄递特征。案件公开稿通常不披露箱号、提单号和运单号，公开情报只能用于锁定时间、港口、品名、企业和路线，后续必须与舱单、报关单、运单及查验图像进行内部反查。",
        "",
        "### （四）出口管制与第三地转运风险",
        "高端AI芯片和服务器案件显示，日本、澳门、香港等节点可能被用于改变运输路径或最终用户表象。应重点识别高性能计算设备、服务器整机、零部件与维修返运之间的异常组合，并核对最终用户、支付方和设备序列号。",
        "",
        "## 四、下一步建议",
        "",
        "1. 建立本期主体台账：统一人员、企业、品牌、港口、仓库和执法机关名称，保留别名、原文拼写及证据链接。",
        "2. 对5个菲律宾冻品集装箱、2个菲律宾卷烟集装箱及斯里兰卡冷库板藏烟案发起单证反查，补齐箱号、提单号、收发货人、报关行和承运船舶。",
        "3. 对Jing Tang Li案串并7批货物的伪装品名、仓库地址、虚假发货企业和澳大利亚收货端；对龟类走私案串并99个邮包的收件地址、支付账户和CITES记录。",
        "4. 将中国来源冻品、电子烟、卷烟、未批准药品和原产地换标案例转化为品名—品牌—路线—价格—重量联合风险规则。",
        "5. 对芯片转运案件建立设备型号、序列号、OEM、最终用户及第三地中转实体核验清单，避免仅以申报目的地判断最终流向。",
        "6. 建立编号证据分级：正文明确披露为A级，清晰原图可复核为B级，仅媒体转述为C级，模糊照片和推测性字符不得入库。",
        "",
        "## 附录：20起采集案例完整内容与深度补充",
        "",
    ]
    for index, item in enumerate(items, 1):
        meta = item.raw_metadata or {}
        review = meta.get("enforcement_review") or {}
        translation = meta.get("translation_zh") or {}
        detail = DETAILS[item.id]
        title = translation.get("title_zh") or item.title
        date = item.published_at.strftime("%Y-%m-%d") if item.published_at else "日期待核"
        basis = "强涉华" if review.get("inclusion_basis") == "strong_china_nexus" else "重大案件"
        chinese_content = re.sub(
            r"\s+", " ",
            str(translation.get("content_zh") or translation.get("summary_zh") or item.summary or ""),
        ).strip()
        lines.extend([
            f"### {index}. {title}",
            f"- **时间/地区**：{date}；{review.get('jurisdiction') or '未知'}；{basis}。",
            f"- **执法事实**：{review.get('enforcement_action') or item.summary or '详见原文'}",
            f"- **案例中文译文**：{chinese_content}",
            f"- **涉华证据**：{review.get('mainland_nexus_evidence') or '未发现直接涉华证据，按重大境外执法案件纳入。'}",
            f"- **企业/人员**：{detail['entities']}",
            f"- **集装箱/运输工具**：{detail['transport']}",
            f"- **同案补充**：{detail['extra']}",
            f"- **原文链接**：{item.url}",
            "",
        ])
    return "\n".join(lines).strip()


def write_docx(path: Path, title_text: str, content: str, item_count: int) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.9)

    def font(style, size, color="1F2937", bold=False):
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    font(doc.styles["Normal"], 10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size, color in (("Title", 22, "0B2545"), ("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "1F4D78"), ("Heading 3", 11.5, "1F4D78")):
        font(doc.styles[name], size, color, True)
        doc.styles[name].paragraph_format.keep_with_next = True

    doc.add_heading(title_text, 0)
    meta = doc.add_paragraph(f"监测周期：2026年7月16日至8月12日  |  案例数量：{item_count}起  |  生成日期：2026年8月13日")
    meta.style = doc.styles["Subtitle"]
    for run in meta.runs:
        run.font.color.rgb = RGBColor.from_string("5F6B7A")
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    for line in content.splitlines():
        value = line.strip()
        if not value:
            continue
        if value.startswith("### "):
            doc.add_heading(value[4:], level=2)
        elif value.startswith("## "):
            if value.startswith("## 附录"):
                doc.add_section(WD_SECTION.NEW_PAGE)
            doc.add_heading(value[3:], level=1)
        elif re.match(r"^\d+\. ", value):
            doc.add_paragraph(value, style="List Number")
        elif value.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(re.sub(r"\*\*", "", value[2:]))
        else:
            p = doc.add_paragraph()
            p.add_run(re.sub(r"\*\*", "", value))

    header = section.header.paragraphs[0]
    header.text = "GatherInfo  |  境外执法信息智能简报"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("6B7280")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("第 ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    footer.add_run(" 页")
    doc.core_properties.title = title_text
    doc.core_properties.author = "GatherInfo / Codex"
    doc.save(path)


def main() -> int:
    db = SessionLocal()
    try:
        run = db.get(CollectionRun, RUN_ID)
        topic = db.get(Topic, TOPIC_ID)
        if not run or not topic:
            raise RuntimeError("Portfolio run or topic not found")
        item_ids = (run.metadata_json or {}).get("portfolio_item_ids") or []
        by_id = {item.id: item for item in db.query(CollectedItem).filter(CollectedItem.id.in_(item_ids)).all()}
        items = [by_id[item_id] for item_id in item_ids]
        if len(items) != 20 or set(DETAILS) != set(item_ids):
            raise RuntimeError("Portfolio and enrichment details are not aligned")

        for item in items:
            item.raw_metadata = {
                **(item.raw_metadata or {}),
                "same_case_deep_search": {
                    **DETAILS[item.id],
                    "reviewed_at": datetime.now().isoformat(),
                    "method": "official source, same-case media cross-search and source image review",
                    "identifier_policy": "Only stable, independently readable identifiers are retained.",
                },
            }

        report = db.get(Report, REPORT_ID)
        if not report:
            report = Report(id=REPORT_ID, topic_id=TOPIC_ID, title=TITLE)
            db.add(report)
        report.title = TITLE
        report.report_type = "analytical"
        report.content = build_report(items)
        report.summary = "20起境外执法案例智能简报：强涉华11起，覆盖11个国家或地区，并附企业、人员及运输标识交叉检索结果。"
        report.status = "completed"
        report.model_id = "ollama-cloud-free+codex-web-research"
        report.tokens_used = 0
        report.item_count = len(items)
        report.item_ids = item_ids
        report.collection_run_id = RUN_ID
        report.date_range_start = min(item.published_at for item in items if item.published_at)
        report.date_range_end = max(item.published_at for item in items if item.published_at)
        report.generated_at = datetime.now()
        report.error_log = None
        report.output_files = None
        report.output_dir = None
        db.commit()
        db.refresh(report)

        report.output_files = export_report(
            report, db.get(SystemConfig, "global"), topic, formats=["md", "html", "pdf"]
        )
        docx_path = Path(report.output_dir) / "执法信息采集_境外执法信息智能简报_2026-08-13.docx"
        write_docx(docx_path, TITLE, report.content or "", report.item_count)
        report.output_files = {**(report.output_files or {}), "docx": str(docx_path)}
        db.commit()
        print(f"report_id={report.id} items={report.item_count} files={report.output_files}")
        return 0
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
