from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt

from app.database import SessionLocal
from app.models import CollectedItem, Report, Topic
from app.report_export import export_report


TOPIC_ID = "foreign-trade-forum-risk-monitoring"
RUN_ID = "codex-forum-high-customs-risk-20260813"
REPORT_ID = "report-forum-high-customs-risk-20260813"
TITLE = "外贸论坛高海关监管风险线索报告（2026年8月13日）"
NOW = datetime.now(timezone.utc)

PRIORITY = {
    "forum-high-customs-risk-20260813-02": "一级",
    "forum-high-customs-risk-20260813-04": "一级",
    "forum-high-customs-risk-20260813-01": "二级",
    "forum-high-customs-risk-20260813-03": "二级",
    "forum-high-customs-risk-20260813-07": "二级",
}


def content_for(items: list[CollectedItem]) -> str:
    lines = [
        f"# {TITLE}", "",
        "## 一、总体结论", "",
        "本轮论坛直连源未产生合格新增，随后通过定向搜索和原帖复核纳入10条高风险线索。线索集中于DDP低报与第三方进口商、货证及HS编码不一致、原产地证主体错配、食品和锂电池监管文件、第三国换标转运。论坛内容均为待核查陈述，不直接作为违法认定。", "",
        "**优先顺序：** 一级2条，具备明确企业或品牌、货物、金额或单证差异；二级3条，具备中国区域、路线和异常事件；三级5条，主要用于识别货代、拼箱及转运模式。", "",
        "## 二、核心发现", "",
        "1. DDP拼箱是最集中风险场景，常同时出现第三方进口商、无法取得清关单、笼统品名和申报价格不可见。",
        "2. 山东聊城钢制螺旋桩线索的企业、实际货物、提单品名及两个HS编码均较明确，最接近可闭环核查。",
        "3. OneXPlayer订单线索给出3107美元成交金额与299美元建议申报值，适合用订单、支付和快件申报数据交叉验证。",
        "4. 原产地证出口人变更、肽类换标改道和食品拆单，应分别联查代理关系、生产批号及检疫文件。", "",
        "## 三、优先核查清单", "",
    ]
    for index, item in enumerate(items, 1):
        priority = PRIORITY.get(item.id, "三级")
        parts = (item.content or "").split("\n\n")
        lines.extend([
            f"### {index}. {item.title}", "",
            f"- 核查等级：{priority}",
            f"- 发布时间：{item.published_at.date().isoformat() if item.published_at else '待核'}",
            f"- {parts[0] if parts else item.summary}",
            f"- {parts[1] if len(parts) > 1 else '海关监管风险：待进一步研判。'}",
            f"- {parts[2] if len(parts) > 2 else '后续核查：调取原始单证复核。'}",
            f"- 原帖：{item.url}", "",
        ])
    lines.extend([
        "## 四、共性排查建议", "",
        "- 以企业、品牌、订单号、支付金额、发货区域、目的港和时间窗串联订单、商业发票、出口报关单、舱单、主分提单及境外清关单。",
        "- 对DDP拼箱统计同一货代、境外进口商和集装箱内多货主申报情况，识别笼统品名、单位价格异常、分票与总票无法勾稽。",
        "- 对HS归类异常比对历史申报、商品规格、材质、用途及更改单；对锂电池联查UN38.3、MSDS和危险品订舱资料。",
        "- 对原产地和第三国转运核对生产商、出口代理、原产地证申请主体、两程提单、换标换箱及实质性加工记录。",
        "- 只有论坛陈述与报关、物流、实物、资金或境外官方单证相互印证后，才升级为确定性风险结论。", "",
        "## 五、证据边界", "",
        "本报告用于风险筛查。论坛帖子及评论属于用户生成内容，可能存在误述、缺漏或营销性表达；涉及企业和品牌的内容均应通过平台订单、企业登记、海关申报及原始物流文件进一步核实。",
    ])
    return "\n".join(lines)


def write_docx(path: Path, content: str) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"].font.size = Pt(10.5)
    for line in content.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif len(line) > 3 and line[0].isdigit() and line[1:3] == ". ":
            doc.add_paragraph(line[3:], style="List Number")
        else:
            doc.add_paragraph(line.replace("**", ""))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def main() -> None:
    db = SessionLocal()
    try:
        topic = db.get(Topic, TOPIC_ID)
        items = db.query(CollectedItem).filter(CollectedItem.run_id == RUN_ID).order_by(CollectedItem.published_at.desc()).all()
        if not topic or len(items) != 10:
            raise RuntimeError(f"expected 10 items, got {len(items)}")
        content = content_for(items)
        report = db.get(Report, REPORT_ID) or Report(id=REPORT_ID, topic_id=TOPIC_ID, title=TITLE)
        report.title, report.report_type, report.content = TITLE, "analytical", content
        report.summary = "复核10条论坛高风险线索，重点关注山东聊城钢制螺旋桩货证不符、OneXPlayer建议低申报值、DDP拼箱第三方进口商及原产地证主体错配。"
        report.status, report.model_id = "completed", "codex-directed-review"
        report.item_count, report.item_ids = len(items), [item.id for item in items]
        report.collection_run_id = RUN_ID
        report.date_range_start = min(item.published_at for item in items if item.published_at)
        report.date_range_end = max(item.published_at for item in items if item.published_at)
        report.generated_at = NOW
        db.add(report)
        db.commit()
        files = export_report(report, {}, topic, formats=["md", "html", "pdf"])
        out_dir = Path(report.output_dir or Path(next(iter(files.values()))).parent)
        docx = out_dir / f"{TITLE}.docx"
        write_docx(docx, content)
        report.output_files = {**files, "docx": str(docx)}
        db.commit()
        print({"report_id": REPORT_ID, "items": len(items), "files": report.output_files})
    finally:
        db.close()


if __name__ == "__main__":
    main()
