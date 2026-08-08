"""Build, review, and persist the three-round customs hotspot report.

The script intentionally separates document generation from database writes:

    --build-only  create the DOCX for visual review, without changing the DB
    --commit      persist the reviewed item/report and switch future exports to DOCX
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_DIR = Path(__file__).resolve().parents[2]
WORKSPACE_DIR = PROJECT_DIR.parent
DB_PATH = PROJECT_DIR / "data" / "gather.db"
OUTPUT_DIR = WORKSPACE_DIR / "outputs" / "时政热点信息分析"
OUTPUT_PATH = OUTPUT_DIR / "涉进出口时政热点三轮优化审核报告_2026-07-16至2026-08-05.docx"
BACKUP_DIR = WORKSPACE_DIR / "work" / "political_hotspots" / "backups"

TOPIC_ID = "weekly-trade-current-affairs"
SOURCE_ID = "ai-smart-web-research"
RUN_ID = "run-customs-three-round-review-20260805"
REPORT_ID = "rpt-customs-hotspots-three-round-reviewed-20260805"
REPORT_TITLE = "涉进出口时政热点三轮优化审核报告（2026年7月16日至8月5日）"
WINDOW_START = "2026-07-16 00:00:00"
WINDOW_END = "2026-08-05 23:59:59"
GENERATED_AT = "2026-08-05 16:30:00"

NAVY = "16324F"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202833"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "F4F7FA"
GREEN = "1F6B4A"
RED = "9B1C1C"
GOLD = "7A5A00"


NEW_ITEM = {
    "id": "customs-hotspot-three-round-20260805-01",
    "title": "中国LNG恢复保税转口，价差驱动下需核查保税账册与复出口舱单一致性（高风险）",
    "published_at": "2026-07-22 00:00:00",
    "category": "保税能源与转口监管",
    "url": "https://www.spglobal.com/energy/en/news-research/latest-news/lng/072226-china-restarts-lng-re-exports-on-sufficient-summer-supply-stronger-jkm",
    "source_name": "S&P Global Commodity Insights；Reuters（公开转载）",
    "source_type": "行业数据报道与新闻报道",
    "facts": (
        "S&P Global 7月22日报道，中国买方在暂停约三个月后恢复商业性LNG复出口。"
        "Maran Gas Posidonia轮7月2日在中海油滨海LNG接收站装载约67695吨，7月10日在日本知多卸货；"
        "Mu Lan轮7月15日在海南装载并于7月22日离港，市场信息指向泰国。同期中国6月LNG进口量环比增长16.8%至568万吨。"
        "Reuters 7月16日报道，美国Plaquemines项目货物由Al Fat'h轮于7月15日至16日抵达海南洋浦，"
        "洋浦保税储罐允许货物在不进入境内市场、不缴进口关税的情况下储存和复出口。"
    ),
    "transmission_chain": (
        "境外LNG装船 -> 中国洋浦或滨海接收站卸载/保税储存 -> 境内外价差形成复出口激励 -> "
        "中国口岸重新装船 -> 日本、泰国等境外目的地。"
    ),
    "risk": (
        "该链路直接进入中国海关保税监管、港口舱单和复出口申报环节。需重点防范保税货物数量与罐容账册不一致、"
        "合理损耗被放大、保税货物违规内销、同批货物原产地或贸易国信息前后不一致，以及进口、转口和复出口价格异常。"
        "公开报道能够证明复出口活动和价差动机，但不证明相关违规已经发生。"
    ),
    "checks": [
        "调取海南洋浦、中海油滨海接收站相关保税入区单、海关账册、罐容计量记录和出区/复出口申报单。",
        "联查Al Fat'h、Maran Gas Posidonia、Mu Lan轮的IMO、AIS、到离港、进口舱单、出口舱单及提单。",
        "核对卸载量、库存量、重新装载量、气化和自然损耗率，识别账实差、异常损耗及保税货物内销未补税。",
        "比对原产国、启运港、贸易国、最终目的港、成交单价和运保费，关注贸易主体或目的地临时变更。",
        "将S&P披露的境内外价差及每百万英热单位超过3美元的理论转口利润作为价格异常筛查基准之一。",
    ],
    "related_sources": [
        {
            "label": "Reuters转载：首票美国LNG抵达洋浦并可能复出口",
            "url": "https://www.sahmcapital.com/news/content/first-us-lng-cargo-since-tariff-dispute-reaches-china-may-be-re-exported-2026-07-16",
        },
    ],
}


AUDIT = {
    NEW_ITEM["id"]: {
        "score": 98,
        "decision": "纳入",
        "stage": "保税监管；中国港口与舱单；复出口申报",
        "reason": "存在实际中国接收站、具体船货、保税储罐和复出口流向，可按账册、舱单、计量及量价闭环核查。",
        "foreign_enforcement_only": False,
    },
    "codex-high-value-hotspot-20260805-01": {
        "score": 97,
        "decision": "纳入",
        "stage": "中国出口管制；许可证与最终用户审核",
        "reason": "中国对无人机及关键部件实施逐案审查，受控商品、许可证、最终用户和目的国均可核验。",
        "foreign_enforcement_only": False,
    },
    "codex-customs-risk-hotspot-20260804-06": {
        "score": 95,
        "decision": "纳入",
        "stage": "中国出口管制；境外最终用户穿透",
        "reason": "14家受限实体和中国来源两用物项范围明确，可直接核查代采主体、最终用户及用途。",
        "foreign_enforcement_only": False,
    },
    "codex-customs-risk-hotspot-20260804-01": {
        "score": 93,
        "decision": "纳入",
        "stage": "中俄陆路口岸；边境反走私；商品归类",
        "reason": "俄境内燃油短缺叠加陆路邻接和价差动机，可核查边境车辆、油箱、HS 2710量价及品名迁移。",
        "foreign_enforcement_only": False,
    },
    "codex-deep-hotspot-20260804-01": {
        "score": 92,
        "decision": "纳入",
        "stage": "中国进口舱单；原产地与完税价格审核",
        "reason": "报道披露驶往舟山的具体中国目的港油轮，可按IMO、AIS、装货港、运保费和舱单核查。",
        "foreign_enforcement_only": False,
    },
    "codex-high-value-hotspot-20260805-04": {
        "score": 89,
        "decision": "纳入",
        "stage": "中哈管道及铁路进口；原产地与价格审核",
        "reason": "CPC中断造成可量化的减产和改道压力，中哈管道、铁路罐车、油种指标及成交价格具备核验条件。",
        "foreign_enforcement_only": False,
    },
    "codex-high-value-hotspot-20260805-03": {
        "score": 87,
        "decision": "纳入",
        "stage": "中国进口舱单；检验检疫；原产地审核",
        "reason": "黑海粮食化肥运输受扰具有供应依赖和航线传导基础，可核查改港换船、原产地、品质和价格。",
        "foreign_enforcement_only": False,
    },
    "codex-deep-hotspot-20260804-05": {
        "score": 85,
        "decision": "纳入",
        "stage": "中国战略矿产出口管制；成分与许可证审核",
        "reason": "中国稀土出口限制及境外短缺直接形成规避动机，商品成分、牌号、许可证和第三国流向可核查。",
        "foreign_enforcement_only": False,
    },
    "codex-high-value-hotspot-20260805-02": {
        "score": 61,
        "decision": "剔除",
        "stage": "美国进口执法",
        "reason": "UFLPA清单主要服务美国进口拦截；当前材料未形成中国海关可直接采取的监管动作。",
        "foreign_enforcement_only": True,
    },
    "5d3dcc660562a378": {
        "score": 48,
        "decision": "剔除",
        "stage": "美国出口管制",
        "reason": "美国关键矿物废料出口授权尚未细化，且直接监管主体为美国海关或商务部门。",
        "foreign_enforcement_only": True,
    },
    "codex-customs-risk-hotspot-20260804-09": {
        "score": 52,
        "decision": "剔除",
        "stage": "待证明",
        "reason": "区域大米价格波动未建立具体中国进口路线、异常主体或检疫措施变化，推演区分度不足。",
        "foreign_enforcement_only": False,
    },
    "codex-customs-risk-hotspot-20260804-04": {
        "score": 50,
        "decision": "剔除",
        "stage": "印度进口政策",
        "reason": "印度国内降税与短缺尚未形成中国进出口货流或中国口岸风险的客观传导证据。",
        "foreign_enforcement_only": True,
    },
    "codex-deep-hotspot-20260804-02": {
        "score": 57,
        "decision": "剔除",
        "stage": "越南进出口监管",
        "reason": "越南强迫劳动产品禁令主要落在越南及目的国合规环节，对中国海关的直接监管动作不明确。",
        "foreign_enforcement_only": True,
    },
    "codex-deep-hotspot-20260804-04": {
        "score": 55,
        "decision": "剔除",
        "stage": "美国关税执法",
        "reason": "关税调整可能影响企业供应链，但海关执法动作主要发生在美国进口端。",
        "foreign_enforcement_only": True,
    },
    "codex-customs-risk-hotspot-20260804-03": {
        "score": 64,
        "decision": "剔除",
        "stage": "待证明",
        "reason": "印度港口拥堵和DAP偏紧与中国海关之间缺少已发生的具体货流、价差或转口路径证据。",
        "foreign_enforcement_only": False,
    },
    "codex-customs-risk-hotspot-20260804-07": {
        "score": 47,
        "decision": "剔除",
        "stage": "美国国防采购与进口合规",
        "reason": "美国国防供应链豁免调整属于境外国防采购规则，中国海关缺少直接处置节点。",
        "foreign_enforcement_only": True,
    },
    "codex-deep-hotspot-20260804-03": {
        "score": 43,
        "decision": "剔除",
        "stage": "美国反规避调查",
        "reason": "美国对埃塞俄比亚、越南光伏产品的反规避调查属于美国进口贸易救济，中国海关风险较小。",
        "foreign_enforcement_only": True,
    },
}


SELECTED_IDS = [
    NEW_ITEM["id"],
    "codex-high-value-hotspot-20260805-01",
    "codex-customs-risk-hotspot-20260804-06",
    "codex-customs-risk-hotspot-20260804-01",
    "codex-deep-hotspot-20260804-01",
    "codex-high-value-hotspot-20260805-04",
    "codex-high-value-hotspot-20260805-03",
    "codex-deep-hotspot-20260804-05",
]


def json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def clean_text(value: object) -> str:
    return " ".join(str(value or "").split())


def parse_content(content: str | None) -> dict[str, str]:
    text = content or ""
    aliases = {
        "source": ["信息来源"],
        "facts": ["原始事件事实", "事实依据"],
        "risk": ["监管风险推演", "监管风险"],
        "checks": ["建议数据核查", "数据核查"],
    }
    result: dict[str, str] = {}
    all_labels = [label for labels in aliases.values() for label in labels]
    for key, labels in aliases.items():
        for label in labels:
            match = re.search(
                rf"(?:^|\n){re.escape(label)}[：:]\s*(.+?)(?=\n\n(?:{'|'.join(map(re.escape, all_labels))})[：:]|\Z)",
                text,
                re.DOTALL,
            )
            if match:
                result[key] = clean_text(match.group(1))
                break
    return result


def load_records() -> tuple[dict[str, dict], list[dict]]:
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    try:
        rows = con.execute(
            """
            SELECT id,title,content,summary,url,category,published_at,raw_metadata
            FROM collected_items
            WHERE topic_id=? AND published_at>=? AND published_at<=?
            ORDER BY published_at DESC, title
            """,
            (TOPIC_ID, WINDOW_START, WINDOW_END),
        ).fetchall()
    finally:
        con.close()

    records: dict[str, dict] = {}
    for row in rows:
        metadata = json.loads(row["raw_metadata"] or "{}")
        parts = parse_content(row["content"])
        records[row["id"]] = {
            "id": row["id"],
            "title": row["title"],
            "content": row["content"] or "",
            "summary": row["summary"] or "",
            "url": row["url"] or "",
            "category": row["category"] or "未分类",
            "published_at": str(row["published_at"] or "")[:10],
            "source_name": metadata.get("source_name") or metadata.get("source_note") or parts.get("source") or urlparse(row["url"] or "").netloc,
            "source_type": metadata.get("source_type") or "公开报道/公告",
            "facts": parts.get("facts") or row["summary"] or "",
            "risk": parts.get("risk") or row["summary"] or "",
            "checks_text": parts.get("checks") or "",
        }

    records[NEW_ITEM["id"]] = {
        **NEW_ITEM,
        "published_at": NEW_ITEM["published_at"][:10],
        "checks_text": "；".join(NEW_ITEM["checks"]),
    }
    missing = set(AUDIT) - set(records)
    if missing:
        raise RuntimeError(f"近20日审核条目缺失：{sorted(missing)}")
    audited = [records[item_id] for item_id in AUDIT]
    return records, audited


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent=120) -> None:
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run_font(run, size=11, bold=None, color=INK, italic=None, east_asia="微软雅黑") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "微软雅黑")
    props.append(fonts)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.append(begin)
    field_run.append(instr)
    field_run.append(separate)
    field_run.append(value)
    field_run.append(end)
    paragraph._p.append(field_run)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.38)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    style_specs = {
        "Title": (23, NAVY, 0, 7),
        "Subtitle": (12, MUTED, 0, 12),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in style_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.05

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1

    if "Meta" not in doc.styles:
        meta = doc.styles.add_style("Meta", WD_STYLE_TYPE.PARAGRAPH)
        meta.font.name = "Calibri"
        meta._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        meta.font.size = Pt(10)
        meta.font.color.rgb = RGBColor.from_string(MUTED)
        meta.paragraph_format.space_after = Pt(3)
        meta.paragraph_format.line_spacing = 1.05

    if "Lead" not in doc.styles:
        lead = doc.styles.add_style("Lead", WD_STYLE_TYPE.PARAGRAPH)
        lead.font.name = "Calibri"
        lead._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        lead.font.size = Pt(11.5)
        lead.font.bold = True
        lead.font.color.rgb = RGBColor.from_string(NAVY)
        lead.paragraph_format.space_before = Pt(4)
        lead.paragraph_format.space_after = Pt(10)
        lead.paragraph_format.left_indent = Inches(0.16)
        lead.paragraph_format.right_indent = Inches(0.16)
        lead.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.text = "涉进出口时政热点  |  三轮优化审核"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(3)
    set_run_font(header.runs[0], 9, bold=True, color=MUTED)
    p_pr = header._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D7DEE8")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer.paragraphs[0]
    add_page_field(footer)


def shade_paragraph(paragraph, fill=PALE_BLUE, border=LIGHT_BLUE) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "6")
        node.set(qn("w:color"), border)
        p_bdr.append(node)
    p_pr.append(p_bdr)


def add_label_paragraph(doc: Document, label: str, value: str, *, color=INK) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    lead = p.add_run(f"{label}：")
    set_run_font(lead, 11, bold=True, color=DARK_BLUE)
    body = p.add_run(value)
    set_run_font(body, 11, color=color)


def add_source_paragraph(doc: Document, record: dict) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    lead = p.add_run("原文链接：")
    set_run_font(lead, 10.5, bold=True, color=DARK_BLUE)
    add_hyperlink(p, record["url"], record["url"])
    for source in record.get("related_sources", []):
        p.add_run("；")
        add_hyperlink(p, source["label"], source["url"])


def build_report_content(records: dict[str, dict], audited: list[dict]) -> str:
    selected = [records[item_id] for item_id in SELECTED_IDS]
    excluded = sorted(
        (record for record in audited if AUDIT[record["id"]]["decision"] == "剔除"),
        key=lambda record: (-AUDIT[record["id"]]["score"], record["published_at"]),
    )
    lines = [
        f"# {REPORT_TITLE}",
        "",
        "## 一、综合分析研判",
        "",
        "近20日审核17条候选信息，按中国海关监管落点、客观传导证据和可执行核查方向三项硬门槛，最终纳入8条、剔除9条。三轮新增检索仅形成1条可正式入库的新线索，即中国LNG保税转口恢复；其余宽泛供需冲击、境外贸易救济和发布日期不可核验信息均未入库。",
        "",
        "风险主要集中于保税能源转口、无人机及两用物项出口管制、陆路边境成品油价差、能源船舶改道、黑海粮食化肥运输和稀土出口管制。建议将公开源线索转化为保税账册、舱单、许可证、最终用户、AIS、量价、商品成分和车辆轨迹等内部核查任务。",
        "",
        "## 二、审核结论",
    ]
    for index, record in enumerate(selected, 1):
        audit = AUDIT[record["id"]]
        lines.append(f"{index}. {record['title']}（{audit['score']}分；{audit['stage']}）")
    lines.extend(["", "## 三、剔除条目及原因"])
    for record in excluded:
        audit = AUDIT[record["id"]]
        lines.append(f"- {record['title']}：{audit['reason']}")
    lines.extend(["", "## 四、入选信息条目"])
    for index, record in enumerate(selected, 1):
        audit = AUDIT[record["id"]]
        lines.extend([
            "",
            f"### {index}. {record['title']}",
            f"- 发布时间：{record['published_at']}；分类：{record['category']}；审核评分：{audit['score']}分。",
            f"- 海关监管落点：{audit['stage']}。",
            f"- 事实依据：{record.get('facts','')}",
            f"- 监管风险：{record.get('risk','')}",
            f"- 核查方向：{record.get('checks_text','')}",
            f"- 原文链接：{record['url']}",
        ])
    return "\n".join(lines).strip()


def build_docx(records: dict[str, dict], audited: list[dict]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("涉进出口时政热点\n三轮优化审核报告")
    set_run_font(run, 23, bold=True, color=NAVY)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("审核期：2026年7月16日至8月5日  |  形成时间：2026年8月5日")

    for label, value in (
        ("审核范围", "主题近20日16条存量信息 + 三轮新增候选信息"),
        ("审核结论", "17条候选中纳入8条、剔除9条；三轮新增正式入库1条"),
        ("核心口径", "必须落到中国海关监管环节，具有客观传导证据，并给出可执行的数据核查方向"),
        ("输出形式", "Word报告；不生成Markdown版本"),
    ):
        p = doc.add_paragraph(style="Meta")
        lead = p.add_run(f"{label}：")
        set_run_font(lead, 10, bold=True, color=DARK_BLUE)
        body = p.add_run(value)
        set_run_font(body, 10, color=MUTED)

    lead = doc.add_paragraph(style="Lead")
    lead.add_run(
        "审核结论：本轮不以“原文是否出现中国”为准，而以风险能否客观传导至中国进出口、边境、口岸、保税、"
        "出口管制或检验检疫环节为准。美国反规避、UFLPA等纯境外执法信息不再纳入。"
    )
    shade_paragraph(lead)

    doc.add_heading("一、综合分析研判", level=1)
    paragraphs = [
        "近20日共审核17条候选信息，其中16条来自主题存量、1条来自三轮优化采集。按照“中国海关监管落点、客观传导证据、可执行核查方向”三项硬门槛，最终保留8条、剔除9条。严格审核后条目数量减少，但每条均能够进一步转化为报关、舱单、保税账册、许可证、船舶或口岸数据核查任务。",
        "从风险结构看，能源类线索最集中。中国LNG恢复保税转口直接涉及保税储罐、进出区账册、复出口舱单和量价核对；霍尔木兹与红海航线扰动、哈萨克斯坦CPC出口中断和黑海港航风险，则可能引发中国进口货物改港、换船、绕航、价格和原产地信息变化。俄罗斯炼厂持续受袭另形成陆路边境价差风险，需从车辆油箱、成品油品名迁移和边贸企业异常入手。",
        "出口监管方面，无人机及关键部件逐案审查、14家欧洲实体两用物项限制和稀土供应约束，均直接连接中国出口许可、最终用户、最终用途和商品成分审核。高风险模式不是普通的“供应链重构”，而是受限买方通过第三国代采、关联公司承接、整机拆分、样品或维修件名义以及低含量合金等方式弱化真实用途和最终用户识别。",
        "本轮将美国对埃塞俄比亚及越南光伏产品反规避调查、UFLPA清单、美国强迫劳动关税、越南强迫劳动产品禁令等信息剔除。上述事项可能影响中国企业经营，但其直接执法节点在境外海关或贸易主管部门，未形成中国海关能够立即执行的监管动作，不宜占用本主题报告篇幅。",
    ]
    for text in paragraphs:
        doc.add_paragraph(text)

    doc.add_heading("二、三轮采集与规则优化", level=1)
    rounds = [
        ("第一轮：供需冲击宽搜", "围绕燃油、化肥、粮食、油气和地区冲突检索。宽泛的“全球短缺+中国”组合噪声较高，仅将LNG保税转口事件留作复核，未将一般价格上涨直接等同于走私风险。"),
        ("第二轮：实际货流与监管环节", "改按中国目的港、中国接收站、中国出口许可证和陆路口岸检索。S&P与Reuters对LNG事件形成交叉印证，确认具体接收站、船舶、装载量、目的地和保税操作条件，达到入库门槛。"),
        ("第三轮：单证与异常指标", "进一步要求候选提供具体商品、口岸/路线、可调取单证和异常指标。哈萨克斯坦燃油出口禁令虽有官方页面，但发布日期无法可靠核验，故只进入观察池、不作为本轮新增条目。"),
    ]
    for heading, body in rounds:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(heading)
        set_run_font(run, 11, bold=True, color=DARK_BLUE)
        p.add_run("。" + body)

    doc.add_heading("三、审核结果", level=1)
    doc.add_paragraph("入选条目按综合价值排序。评分用于确定核查优先级，不代表违法行为已经发生。", style="Meta")
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    headers = ["序号", "评分", "信息条目", "中国海关监管落点"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, 9.5, bold=True, color=NAVY)
    for rank, item_id in enumerate(SELECTED_IDS, 1):
        record = records[item_id]
        audit = AUDIT[item_id]
        cells = table.add_row().cells
        values = [str(rank), str(audit["score"]), record["title"], audit["stage"]]
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, 9.3, bold=index == 1, color=GREEN if index == 1 else INK)
    set_table_geometry(table, [600, 720, 5160, 2880])

    doc.add_heading("四、下一步核查建议", level=1)
    recommendations = [
        "保税LNG：以洋浦、滨海接收站为节点，核对保税入区、罐容计量、合理损耗、重新装船、复出口舱单和最终目的港，先验证账实与量价闭环。",
        "两用物项：围绕无人机、关键部件、稀土及受限实体建立“商品参数—许可证—最终用户—目的国—关联企业”五字段联查，识别拆分订单和第三国代采。",
        "海运能源：联查IMO、AIS、船旗、装货港、船对船转运、提单签发地、运保费和油品理化指标，识别改港换船与来源信息不一致。",
        "陆路油品：关注黑龙江、内蒙古及中哈口岸频繁往返车辆、异常油箱容积、空载返程、HS 2710项下品名迁移和小额高频边贸申报。",
        "结果闭环：公开源只用于发现风险，不直接定性。命中异常后应由报关、舱单、企业、物流和资金数据交叉验证，再决定是否布控、查验或移交。",
    ]
    for recommendation in recommendations:
        doc.add_paragraph(recommendation, style="List Number")

    excluded = sorted(
        (record for record in audited if AUDIT[record["id"]]["decision"] == "剔除"),
        key=lambda record: (-AUDIT[record["id"]]["score"], record["published_at"]),
    )
    doc.add_heading("五、剔除条目及原因", level=1)
    doc.add_paragraph("以下9条不进入本期报告正文，也不作为自动采集的正样本。", style="Meta")
    for record in excluded:
        audit = AUDIT[record["id"]]
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(record["title"])
        set_run_font(run, 10.5, bold=True, color=RED if audit["foreign_enforcement_only"] else GOLD)
        p.add_run("。" + audit["reason"])

    doc.add_page_break()
    doc.add_heading("六、入选信息条目", level=1)
    doc.add_paragraph(
        "以下内容先列公开事实，再列中国海关风险传导和数据核查方向。风险研判均不代表相关违法行为已经发生。",
        style="Meta",
    )

    for rank, item_id in enumerate(SELECTED_IDS, 1):
        record = records[item_id]
        audit = AUDIT[item_id]
        heading = doc.add_heading(f"{rank}. {record['title']}", level=2)
        heading.paragraph_format.page_break_before = rank > 1

        p = doc.add_paragraph(style="Meta")
        for label, value in (
            ("发布时间", record["published_at"]),
            ("分类", record["category"]),
            ("审核评分", f"{audit['score']}分"),
        ):
            lead_run = p.add_run(f"{label}：")
            set_run_font(lead_run, 9.5, bold=True, color=DARK_BLUE)
            value_run = p.add_run(value + "   ")
            set_run_font(value_run, 9.5, color=MUTED)

        add_label_paragraph(doc, "信息来源及类型", f"{clean_text(record.get('source_name'))}（{record.get('source_type','公开信息')}）")
        add_label_paragraph(doc, "公开事实", record.get("facts", ""))
        if record.get("transmission_chain"):
            add_label_paragraph(doc, "对华传导链", record["transmission_chain"])
        else:
            add_label_paragraph(doc, "中国海关监管落点", audit["stage"])
        add_label_paragraph(doc, "监管风险研判", record.get("risk", ""))

        checks = record.get("checks")
        if not checks:
            checks_text = clean_text(record.get("checks_text"))
            checks = [part.strip("；。 ") for part in re.split(r"[；;]", checks_text) if part.strip()]
        p = doc.add_paragraph()
        lead_run = p.add_run("建议核查：")
        set_run_font(lead_run, 11, bold=True, color=DARK_BLUE)
        for check in checks:
            doc.add_paragraph(check, style="List Bullet")

        boundary = doc.add_paragraph()
        boundary.paragraph_format.space_before = Pt(2)
        boundary.paragraph_format.space_after = Pt(5)
        lead_run = boundary.add_run("审核边界：")
        set_run_font(lead_run, 10.5, bold=True, color=GOLD)
        text_run = boundary.add_run(
            audit["reason"] + " 公开源用于发现待核风险，不替代海关内部数据和执法调查。"
        )
        set_run_font(text_run, 10.5, color=INK)
        add_source_paragraph(doc, record)

    doc.add_heading("附：后续自动采集硬门槛", level=1)
    gates = [
        "原文可以不出现中国，但必须说明风险如何传导至中国进出口、边境、口岸、保税、出口管制、检验检疫或跨境电商监管。",
        "必须存在地理邻接、实际中国货流、具体船货、显著价差、供应依赖或中国管制措施中的至少一项客观基础。",
        "必须给出具体数据或单证、商品或HS范围、异常指标和后续动作；缺少其中任一项不得自动入库。",
        "纯境外反倾销、反规避、UFLPA、境外市场准入或进口商合规信息，若中国海关无直接监管动作，一律剔除。",
        "来源发布日期无法核验、只有论坛转述或同一事件重复报道的，进入观察池，不生成正式条目。",
    ]
    for gate in gates:
        doc.add_paragraph(gate, style="List Number")

    properties = doc.core_properties
    properties.title = REPORT_TITLE
    properties.subject = "涉进出口时政热点三轮采集优化审核"
    properties.author = "GatherInfo"
    properties.keywords = "开源情报, 海关监管, 进出口风险, 供应链穿透"
    doc.save(OUTPUT_PATH)


def validate_docx(path: Path) -> None:
    if not path.is_file() or path.stat().st_size < 20_000:
        raise RuntimeError(f"DOCX不存在或文件过小：{path}")
    doc = Document(path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    required = [
        "三轮优化审核报告",
        "综合分析研判",
        "中国LNG恢复保税转口",
        "美国对埃塞俄比亚及越南光伏产品",
        "后续自动采集硬门槛",
    ]
    missing = [value for value in required if value not in text]
    if missing:
        raise RuntimeError(f"DOCX结构校验失败，缺少：{missing}")
    if len(doc.paragraphs) < 90 or len(doc.tables) != 1:
        raise RuntimeError("DOCX结构校验失败：段落或表格数量异常")


def backup_database() -> Path:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    path = BACKUP_DIR / f"gather-before-three-round-customs-report-{datetime.now():%Y%m%d-%H%M%S}.db"
    source = sqlite3.connect(DB_PATH)
    target = sqlite3.connect(path)
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()
    return path


def upsert_new_item(con: sqlite3.Connection, now: str) -> None:
    content = (
        f"信息来源：{NEW_ITEM['source_name']}（{NEW_ITEM['source_type']}），2026-07-22。\n\n"
        f"事实依据：{NEW_ITEM['facts']}\n\n"
        f"对华传导链：{NEW_ITEM['transmission_chain']}\n\n"
        f"监管风险：{NEW_ITEM['risk']}\n\n"
        f"数据核查：{'；'.join(NEW_ITEM['checks'])}"
    )
    audit = AUDIT[NEW_ITEM["id"]]
    metadata = {
        "provider": "iterative_open_source_review",
        "import_batch": "customs-three-round-review-20260805",
        "source_name": NEW_ITEM["source_name"],
        "source_type": NEW_ITEM["source_type"],
        "related_sources": NEW_ITEM["related_sources"],
        "customs_hotspot_review": {
            "china_customs_score": 98,
            "transmission_evidence_score": 97,
            "executable_check_score": 98,
            "foreign_enforcement_only": False,
            "china_customs_stage": audit["stage"],
            "transmission_chain": NEW_ITEM["transmission_chain"],
            "customs_data_checks": NEW_ITEM["checks"],
            "analysis_disclaimer": "监管风险为基于公开事实的待核研判，不代表违法行为已经发生。",
        },
        "customs_gate_audit_20260805_v2": audit,
    }
    entities = {
        "ports": ["海南洋浦", "中海油滨海LNG接收站"],
        "vessels": ["Al Fat'h", "Maran Gas Posidonia", "Mu Lan"],
        "destinations": ["日本知多", "泰国（市场信息，待确认）"],
        "customs_stage": audit["stage"],
    }
    values = (
        NEW_ITEM["id"], SOURCE_ID, RUN_ID, TOPIC_ID, NEW_ITEM["title"], content,
        sha256_text(NEW_ITEM["url"] + "\n" + content), NEW_ITEM["risk"], NEW_ITEM["url"],
        "zh", NEW_ITEM["category"], json_text(entities), "enriched", 0.98, 0.99,
        NEW_ITEM["published_at"], now, now, json_text(metadata), "public",
    )
    con.execute(
        """
        INSERT INTO collected_items (
            id,source_id,run_id,topic_id,title,content,content_hash,summary,url,language,category,
            entities,status,quality_score,relevance_score,published_at,collected_at,updated_at,
            raw_metadata,authorization_level
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(id) DO UPDATE SET
            source_id=excluded.source_id,run_id=excluded.run_id,topic_id=excluded.topic_id,
            title=excluded.title,content=excluded.content,content_hash=excluded.content_hash,
            summary=excluded.summary,url=excluded.url,language=excluded.language,category=excluded.category,
            entities=excluded.entities,status=excluded.status,quality_score=excluded.quality_score,
            relevance_score=excluded.relevance_score,published_at=excluded.published_at,
            updated_at=excluded.updated_at,raw_metadata=excluded.raw_metadata,
            authorization_level=excluded.authorization_level
        """,
        values,
    )


def update_item_audits(con: sqlite3.Connection, now: str) -> None:
    for item_id, audit in AUDIT.items():
        if item_id == NEW_ITEM["id"]:
            continue
        row = con.execute("SELECT raw_metadata FROM collected_items WHERE id=?", (item_id,)).fetchone()
        if row is None:
            raise RuntimeError(f"待审核条目不存在：{item_id}")
        metadata = json.loads(row[0] or "{}")
        metadata["customs_gate_audit_20260805_v2"] = audit
        con.execute(
            "UPDATE collected_items SET raw_metadata=?, updated_at=? WHERE id=?",
            (json_text(metadata), now, item_id),
        )


def persist_report(records: dict[str, dict], audited: list[dict]) -> Path:
    validate_docx(OUTPUT_PATH)
    backup = backup_database()
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat(sep=" ")
    docx_hash = hashlib.sha256(OUTPUT_PATH.read_bytes()).hexdigest()
    content = build_report_content(records, audited)
    summary = (
        "三轮优化采集后，对近20日17条候选信息执行中国海关监管落点、客观传导证据和可执行核查方向三项硬审核，"
        "最终纳入8条、剔除9条；新增LNG保税转口线索1条。"
    )
    con = sqlite3.connect(DB_PATH)
    try:
        con.execute("PRAGMA foreign_keys=ON")
        con.execute("BEGIN IMMEDIATE")
        upsert_new_item(con, now)
        update_item_audits(con, now)

        run_metadata = {
            "provider": "iterative_open_source_review",
            "rounds": 3,
            "review_window_days": 20,
            "formal_new_items": 1,
            "reviewed_candidate_items": 17,
            "approved_report_items": 8,
            "review_logic": "china_customs_stage + objective_transmission + executable_check",
            "round_notes": [
                "第一轮宽搜后收紧一般供需冲击",
                "第二轮限定中国实际货流与监管环节",
                "第三轮增加单证、异常指标和发布日期硬门槛",
            ],
        }
        con.execute(
            """
            INSERT INTO collection_runs (
                id,source_id,topic_id,status,batch_id,keywords_used,items_found,items_new,
                items_updated,items_failed,started_at,completed_at,duration_ms,window_start,
                window_end,error_log,metadata_json,created_at,progress_events
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(id) DO UPDATE SET
                status=excluded.status,batch_id=excluded.batch_id,keywords_used=excluded.keywords_used,
                items_found=excluded.items_found,items_new=excluded.items_new,
                items_updated=excluded.items_updated,items_failed=excluded.items_failed,
                completed_at=excluded.completed_at,window_start=excluded.window_start,
                window_end=excluded.window_end,error_log=excluded.error_log,
                metadata_json=excluded.metadata_json,progress_events=excluded.progress_events
            """,
            (
                RUN_ID, SOURCE_ID, TOPIC_ID, "completed", "customs-three-round-review-20260805",
                json_text(["中国实际货流", "保税监管", "边境价差", "许可证", "最终用户", "舱单AIS"]),
                1, 1, 0, 0, now, now, 0, WINDOW_START, WINDOW_END,
                json_text([]), json_text(run_metadata), now,
                json_text([{"stage": "completed", "message": "三轮采集、逐条审核及Word报告复核完成"}]),
            ),
        )

        output_files = {"docx": str(OUTPUT_PATH)}
        report_values = (
            REPORT_ID, TOPIC_ID, REPORT_TITLE, content, summary, "completed",
            "verified-open-source-review", 0, len(SELECTED_IDS), json_text(SELECTED_IDS),
            None, RUN_ID, WINDOW_START, WINDOW_END, json_text(output_files), str(OUTPUT_DIR),
            now, now, "analytical",
        )
        con.execute(
            """
            INSERT INTO reports (
                id,topic_id,title,content,summary,status,model_id,tokens_used,item_count,item_ids,
                error_log,collection_run_id,date_range_start,date_range_end,output_files,output_dir,
                generated_at,created_at,report_type
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(id) DO UPDATE SET
                title=excluded.title,content=excluded.content,summary=excluded.summary,
                status=excluded.status,model_id=excluded.model_id,tokens_used=excluded.tokens_used,
                item_count=excluded.item_count,item_ids=excluded.item_ids,error_log=excluded.error_log,
                collection_run_id=excluded.collection_run_id,date_range_start=excluded.date_range_start,
                date_range_end=excluded.date_range_end,output_files=excluded.output_files,
                output_dir=excluded.output_dir,generated_at=excluded.generated_at,
                report_type=excluded.report_type
            """,
            report_values,
        )

        con.execute(
            """
            UPDATE topics SET
                total_items_collected=(SELECT COUNT(*) FROM collected_items WHERE topic_id=?),
                last_run_at=?,last_collection_run_id=?,last_error=NULL,updated_at=?
            WHERE id=?
            """,
            (TOPIC_ID, now, RUN_ID, now, TOPIC_ID),
        )
        con.execute(
            """
            UPDATE source_configs SET
                items_collected=(SELECT COUNT(*) FROM collected_items WHERE source_id=?),
                last_sync_at=?,last_error=NULL,updated_at=?
            WHERE id=?
            """,
            (SOURCE_ID, now, now, SOURCE_ID),
        )
        con.execute(
            "UPDATE system_config SET report_formats=?, updated_at=? WHERE id='global'",
            (json_text(["docx"]), now),
        )
        con.execute(
            "UPDATE reports SET output_files=json_object('docx', ?), output_dir=? WHERE id=?",
            (str(OUTPUT_PATH), str(OUTPUT_DIR), REPORT_ID),
        )
        con.execute(
            "UPDATE reports SET summary=summary || ? WHERE id=?",
            (f" Word文件SHA256：{docx_hash[:16]}。", REPORT_ID),
        )
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()
    return backup


def main() -> int:
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--build-only", action="store_true")
    mode.add_argument("--commit", action="store_true")
    args = parser.parse_args()

    records, audited = load_records()
    if args.build_only:
        build_docx(records, audited)
        validate_docx(OUTPUT_PATH)
        print(json_text({
            "mode": "build-only",
            "docx": str(OUTPUT_PATH),
            "bytes": OUTPUT_PATH.stat().st_size,
            "candidate_count": len(audited),
            "selected_count": len(SELECTED_IDS),
            "database_changed": False,
        }))
        return 0

    backup = persist_report(records, audited)
    print(json_text({
        "mode": "commit",
        "backup": str(backup),
        "report_id": REPORT_ID,
        "run_id": RUN_ID,
        "docx": str(OUTPUT_PATH),
        "selected_count": len(SELECTED_IDS),
        "new_items": 1,
        "future_report_formats": ["docx"],
    }))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
