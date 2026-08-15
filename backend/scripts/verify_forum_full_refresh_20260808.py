"""Read-only acceptance checks for the 2026-08-08 forum refresh."""

from __future__ import annotations

import hashlib
import json
import sqlite3
import zipfile
from pathlib import Path

from docx import Document

from forum_full_refresh_data_20260808 import (
    NEW_KEYWORDS,
    NEW_LEADS,
    OUTPUT_DOCX,
    PROMPT_ID,
    REPORT_ID,
    REPORT_LEADS,
    SOURCE_AUDIT,
    TOPIC_ID,
)


DB_PATH = Path(__file__).resolve().parents[2] / "data" / "gather.db"
PROMPT_MARKER = "[2026-08-08 全源复核与互联网补充规则]"
SEASONAL_PROMPT_MARKER = "[2026-08-08 中秋国庆时令生鲜夹藏与商业代带专项规则]"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    with sqlite3.connect(DB_PATH) as connection:
        connection.row_factory = sqlite3.Row
        report = connection.execute("SELECT * FROM reports WHERE id = ?", (REPORT_ID,)).fetchone()
        topic = connection.execute("SELECT * FROM topics WHERE id = ?", (TOPIC_ID,)).fetchone()
        prompt = connection.execute("SELECT * FROM prompt_templates WHERE id = ?", (PROMPT_ID,)).fetchone()
        placeholders = ",".join("?" for _ in NEW_LEADS)
        items = connection.execute(
            f"SELECT id, topic_id, source_id, status, url, raw_metadata FROM collected_items WHERE id IN ({placeholders})",
            [lead["id"] for lead in NEW_LEADS],
        ).fetchall()
        audit_runs = connection.execute(
            "SELECT id, source_id, status, batch_id, metadata_json FROM collection_runs WHERE batch_id = ?",
            ("codex-forum-risk-full-source-refresh-20260808",),
        ).fetchall()
        new_sources = connection.execute(
            "SELECT id, is_active, is_configured, last_sync_at FROM source_configs WHERE id IN (?, ?, ?)",
            ("web-linktrans-official", "web-welisen-logistics", "web-zerrand-crossborder-runner"),
        ).fetchall()

    assert report is not None and report["status"] == "completed"
    assert report["item_count"] == len(REPORT_LEADS) == 7
    assert set(json.loads(report["item_ids"])) == {lead["id"] for lead in REPORT_LEADS}
    output_files = json.loads(report["output_files"])
    output_path = Path(output_files["docx"])
    assert output_path == Path(OUTPUT_DOCX) and output_path.is_file()

    assert topic is not None and topic["total_items_collected"] == 27
    assert len(json.loads(topic["source_ids"])) == len(SOURCE_AUDIT) == 39
    assert all(keyword in json.loads(topic["keywords"]) for keyword in NEW_KEYWORDS)
    assert "[2026-08-08 全源复核]" in topic["description_prompt"]
    assert prompt is not None and PROMPT_MARKER in prompt["content"]
    assert SEASONAL_PROMPT_MARKER in prompt["content"]

    assert len(items) == len(NEW_LEADS) == 3
    assert all(row["topic_id"] == TOPIC_ID and row["status"] == "enriched" for row in items)
    gates = {json.loads(row["raw_metadata"])["china_entity_gate"] for row in items}
    assert gates == {"conditional", "passed_by_platform_order"}

    assert len(audit_runs) == len(SOURCE_AUDIT) == 39
    assert all(row["status"] == "completed" for row in audit_runs)
    assert len(new_sources) == 3
    assert all(row["is_active"] and row["is_configured"] and row["last_sync_at"] for row in new_sources)

    document = Document(output_path)
    heading_text = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert "一、本轮结果" in heading_text
    assert "信息源逐源复核台账" in heading_text
    assert "中秋、国庆时令生鲜夹藏与商业代带专项预警" in heading_text
    assert "综合研判与执行顺序" in heading_text
    assert len(document.tables) == 2
    assert len(document.tables[0].rows) == len(REPORT_LEADS) + 1
    assert len(document.tables[1].rows) == len(SOURCE_AUDIT) + 1
    with zipfile.ZipFile(output_path) as archive:
        assert archive.testzip() is None

    print(
        json.dumps(
            {
                "report_status": report["status"],
                "report_items": report["item_count"],
                "topic_items": topic["total_items_collected"],
                "topic_sources": len(json.loads(topic["source_ids"])),
                "new_items": len(items),
                "audit_runs": len(audit_runs),
                "prompt_marker": True,
                "keywords_added": len(NEW_KEYWORDS),
                "docx_tables": len(document.tables),
                "docx_size": output_path.stat().st_size,
                "docx_sha256": sha256(output_path),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
