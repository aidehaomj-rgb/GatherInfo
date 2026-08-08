"""Build and structurally audit the Word report for the 2026-08-08 forum refresh."""

from __future__ import annotations

import os
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_china_entity_recollection_docx_20260805 as base
from forum_full_refresh_data_20260808 import (
    DATE_RANGE_END,
    DATE_RANGE_START,
    OUTPUT_DIR,
    OUTPUT_DOCX,
    REPORT_LEADS,
    REPORT_TITLE,
    SEASONAL_ALERT,
    SOURCE_AUDIT,
)


USABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
SOURCE_LOOKUP = {item["id"]: item for item in SOURCE_AUDIT}


def configure_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("外贸论坛违规清关与走私风险线索")
    base.set_run_font(left, 8.5, base.MUTED, bold=True)
    right = paragraph.add_run("\t全源复核与互联网补充")
    base.set_run_font(right, 8.5, base.MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("公开网络待核线索｜不构成违法认定")
    base.set_run_font(left, 8.5, base.MUTED)
    tab = paragraph.add_run("\t第 ")
    base.set_run_font(tab, 8.5, base.MUTED)
    base.add_field(paragraph, " PAGE ")
    middle = paragraph.add_run(" 页 / 共 ")
    base.set_run_font(middle, 8.5, base.MUTED)
    base.add_field(paragraph, " NUMPAGES ", "10")
    end = paragraph.add_run(" 页")
    base.set_run_font(end, 8.5, base.MUTED)


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph(style="Kicker")
    kicker.add_run("风险线索专题报告")
    title = doc.add_paragraph(REPORT_TITLE, style="Title")
    title.paragraph_format.keep_with_next = True
    subtitle = doc.add_paragraph(
        "近30天｜非执法案例｜中国可核实体优先｜逐源复核后互联网补充",
        style="Subtitle",
    )
    metadata = [
        ("报告主题：", "外贸论坛违规清关与走私风险线索"),
        ("采集窗口：", f"{DATE_RANGE_START}—{DATE_RANGE_END}"),
        ("覆盖范围：", f"{len(SOURCE_AUDIT)}个已配置/新增来源；7组当前核查条目"),
        ("生成日期：", "2026-08-08"),
        ("筛选口径：", "强风险组合 + 中国企业/平台订单/口岸或运输标识 + 同源去重"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph(style="Metadata")
        label_run = paragraph.add_run(label)
        base.set_run_font(label_run, 10.2, base.INK_BLUE, bold=True)
        value_run = paragraph.add_run(value)
        base.set_run_font(value_run, 10.2, base.MUTED)


def add_callout(doc: Document, label: str, text: str, color: str = base.INK_BLUE) -> None:
    paragraph = doc.add_paragraph(style="Risk Note")
    base.add_paragraph_shading(paragraph, base.CALLOUT)
    base.set_paragraph_indents(paragraph, left=160, right=160)
    label_run = paragraph.add_run(label)
    base.set_run_font(label_run, 10.5, color, bold=True)
    text_run = paragraph.add_run(text)
    base.set_run_font(text_run, 10.5, base.INK_BLUE)


def add_summary_table(doc: Document) -> None:
    headers = ["序号", "日期 / 等级", "中国可核实体或标识", "主要风险与首要核查门槛"]
    widths = [600, 1650, 2810, 4300]
    table = doc.add_table(rows=1, cols=4)
    table.style = None
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for column, text in enumerate(headers):
        cell = table.rows[0].cells[column]
        base.shade_cell(cell, base.LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.style = doc.styles["Table Body"]
        run = paragraph.add_run(text)
        base.set_run_font(run, 9.2, base.INK_BLUE, bold=True)
    base.set_repeat_header(table.rows[0])

    for index, lead in enumerate(REPORT_LEADS, 1):
        row = table.add_row()
        first_entity = lead["china_entities"][0]
        values = [
            f"{index:02d}",
            lead["published_at"][:10] + "\n" + lead["risk_level"],
            first_entity,
            "；".join(lead["risk_signals"][:3]) + "。首要门槛：" + lead["next_steps"][0],
        ]
        for column, value in enumerate(values):
            paragraph = row.cells[column].paragraphs[0]
            paragraph.style = doc.styles["Table Body"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column < 2 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            base.set_run_font(run, 9.0)
    base.set_table_geometry(table, widths)
    base.set_table_borders(table)


def add_source_line(doc: Document, lead: dict) -> None:
    source = SOURCE_LOOKUP[lead["source_id"]]
    paragraph = doc.add_paragraph(style="Source Citation")
    label = paragraph.add_run("原帖来源：")
    base.set_run_font(label, 9, base.MUTED, bold=True)
    base.add_hyperlink(paragraph, source["name"], lead["url"])
    if lead.get("crosspost_url"):
        cross = paragraph.add_run("  ｜  跨版链接（已去重）：")
        base.set_run_font(cross, 9, base.MUTED)
        base.add_hyperlink(paragraph, "r/freightforwarding", lead["crosspost_url"])


def add_labeled_paragraph(doc: Document, label: str, text: str, style: str = "Normal") -> None:
    paragraph = doc.add_paragraph(style=style)
    label_run = paragraph.add_run(label)
    base.set_run_font(label_run, 10.5 if style != "Small Note" else 9, base.INK_BLUE, bold=True)
    text_run = paragraph.add_run(text)
    base.set_run_font(text_run, 10.5 if style != "Small Note" else 9, base.MUTED if style == "Small Note" else "000000")


def add_lead_section(doc: Document, lead: dict, index: int, bullet_abstract: int, decimal_abstract: int) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run(f"{index}. {lead['title']}")
    add_source_line(doc, lead)
    add_labeled_paragraph(
        doc,
        "发布时间 / 核查等级：",
        f"{lead['published_at'][:10]}  ｜  {lead['risk_level']}",
        style="Metadata",
    )

    doc.add_heading("帖子主要内容", level=2)
    paragraph = doc.add_paragraph(lead["content_summary"])
    paragraph.paragraph_format.keep_together = True

    doc.add_heading("中国可核实体与运输标识", level=2)
    num_id = base.new_num_id(doc, bullet_abstract)
    for value in lead["china_entities"]:
        base.add_list_paragraph(doc, value, num_id)

    doc.add_heading("风险分析与证据边界", level=2)
    num_id = base.new_num_id(doc, bullet_abstract)
    for value in lead["risk_points"]:
        base.add_list_paragraph(doc, value, num_id)

    doc.add_heading("下一步核查建议", level=2)
    num_id = base.new_num_id(doc, decimal_abstract)
    for value in lead["next_steps"]:
        base.add_list_paragraph(doc, value, num_id)

    note = doc.add_paragraph(style="Risk Note")
    base.add_paragraph_shading(note, "FFF4E5")
    base.set_paragraph_indents(note, left=160, right=160)
    label = note.add_run("当前缺失字段：")
    base.set_run_font(label, 10.5, base.RISK_RED, bold=True)
    body = note.add_run(lead["missing_fields"])
    base.set_run_font(body, 10.5, base.INK_BLUE)

    reliability = doc.add_paragraph(style="Small Note")
    label = reliability.add_run("来源与可靠性说明：")
    base.set_run_font(label, 9, base.CAUTION, bold=True)
    text = reliability.add_run(lead["source_note"])
    base.set_run_font(text, 9, base.MUTED)


def add_source_audit_table(doc: Document) -> None:
    doc.add_heading("信息源逐源复核台账", level=1)
    paragraph = doc.add_paragraph(
        "以下台账记录本轮对全部来源的结果判断。“未保留”不等于来源无价值，只表示本轮未满足近30天、中国可核对象和强风险组合三项条件。"
    )
    paragraph.paragraph_format.keep_with_next = True

    headers = ["来源", "本轮状态", "复核结论"]
    widths = [2540, 1480, 5340]
    table = doc.add_table(rows=1, cols=3)
    table.style = None
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for column, text in enumerate(headers):
        cell = table.rows[0].cells[column]
        base.shade_cell(cell, base.LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.style = doc.styles["Table Body"]
        run = paragraph.add_run(text)
        base.set_run_font(run, 9.2, base.INK_BLUE, bold=True)
    base.set_repeat_header(table.rows[0])

    for audit in SOURCE_AUDIT:
        row = table.add_row()
        source_paragraph = row.cells[0].paragraphs[0]
        source_paragraph.style = doc.styles["Table Body"]
        base.add_hyperlink(source_paragraph, audit["name"], audit["url"])
        for column, value in ((1, audit["status"]), (2, audit["note"])):
            paragraph = row.cells[column].paragraphs[0]
            paragraph.style = doc.styles["Table Body"]
            run = paragraph.add_run(value)
            base.set_run_font(run, 8.8)
    base.set_table_geometry(table, widths)
    base.set_table_borders(table)


def add_seasonal_alert_section(doc: Document, bullet_abstract: int, decimal_abstract: int) -> None:
    doc.add_heading("中秋、国庆时令生鲜夹藏与商业代带专项预警", level=1)
    add_callout(
        doc,
        "本轮判断：",
        "未发现近30天内同时公开夹藏方法、具体数量和中国经营主体的合格帖子，因此不新增当期风险条目。以下内容属于节前预警和超窗模式样本。",
        color=base.CAUTION,
    )

    heading = doc.add_paragraph(style="Heading 2")
    heading.add_run(SEASONAL_ALERT["title"])

    paragraph = doc.add_paragraph(style="Source Citation")
    label = paragraph.add_run("公开页面：")
    base.set_run_font(label, 9, base.MUTED, bold=True)
    base.add_hyperlink(paragraph, "当前生鲜跨境服务页", SEASONAL_ALERT["service_url"])
    separator = paragraph.add_run("  ｜  ")
    base.set_run_font(separator, 9, base.MUTED)
    base.add_hyperlink(paragraph, "历史大闸蟹案例页", SEASONAL_ALERT["case_url"])

    doc.add_heading("页面内容与季节性信号", level=2)
    doc.add_paragraph(SEASONAL_ALERT["summary"])

    doc.add_heading("中国可核实体与节点", level=2)
    num_id = base.new_num_id(doc, bullet_abstract)
    for value in SEASONAL_ALERT["entities"]:
        base.add_list_paragraph(doc, value, num_id)

    doc.add_heading("风险信号", level=2)
    num_id = base.new_num_id(doc, bullet_abstract)
    for value in SEASONAL_ALERT["risk_signals"]:
        base.add_list_paragraph(doc, value, num_id)

    doc.add_heading("证据边界与误判控制", level=2)
    num_id = base.new_num_id(doc, bullet_abstract)
    for value in SEASONAL_ALERT["boundary"]:
        base.add_list_paragraph(doc, value, num_id)

    official = doc.add_paragraph(style="Source Citation")
    label = official.add_run("官方合规依据：")
    base.set_run_font(label, 9, base.MUTED, bold=True)
    base.add_hyperlink(official, "香港海关旅客携带食物问答", SEASONAL_ALERT["customs_faq_url"])
    official.add_run("  ｜  ")
    base.add_hyperlink(official, "香港食安中心大闸蟹要求", SEASONAL_ALERT["cfs_url"])
    official.add_run("  ｜  ")
    base.add_hyperlink(official, "内地大闸蟹恢复直接供港", SEASONAL_ALERT["direct_supply_url"])

    doc.add_heading("下一步核查建议", level=2)
    num_id = base.new_num_id(doc, decimal_abstract)
    for value in SEASONAL_ALERT["next_steps"]:
        base.add_list_paragraph(doc, value, num_id)

    note = doc.add_paragraph(style="Risk Note")
    base.add_paragraph_shading(note, "FFF4E5")
    base.set_paragraph_indents(note, left=160, right=160)
    label = note.add_run("当前缺失字段：")
    base.set_run_font(label, 10.5, base.RISK_RED, bold=True)
    body = note.add_run(SEASONAL_ALERT["missing_fields"])
    base.set_run_font(body, 10.5, base.INK_BLUE)


def add_method_and_conclusion(doc: Document, bullet_abstract: int, decimal_abstract: int) -> None:
    doc.add_heading("检索逻辑优化与下一轮规则", level=1)
    rules = [
        "原帖时间优先：首次发布时间必须在30天内；抓取时间、页面更新时间或目录刷新时间不得替代。",
        "强风险组合优先：低报/改品名/无证件/查验规避等表述，必须再叠加具体货物、路线、企业、平台订单或运输标识。",
        "中国核查对象硬门槛：境外个人姓名不入选；可通过Alibaba订单、店铺、收款主体、境内口岸/仓库或运输单证回溯到中国主体的可保留。",
        "证据分层：原帖自述、评论推断、企业官网自述和第三方目录信息分别标注，官网仅用于主体消歧。",
        "跨版去重：同一货量、路线和文本的帖子只保留一条主记录，并保存转载链接。",
        "超窗模式库：超出30天但包含价值低申、集中申报、改品名等高风险词的页面只进入监测规则，不进入当期风险清单。",
    ]
    num_id = base.new_num_id(doc, bullet_abstract)
    for rule in rules:
        base.add_list_paragraph(doc, rule, num_id)

    doc.add_heading("综合研判与执行顺序", level=1)
    add_callout(
        doc,
        "结论：",
        "本轮新增3组、合并既有4组，共7组当前核查条目。优先级最高的是带Alibaba订单号的钢制螺旋桩单证差异；其次是Linktrans同名货代低报询问和两组既有中国企业公开业务条目。所有线索均需先锁定主体和原始单证。",
    )
    actions = [
        "第一优先：以Alibaba订单号290841751001029390调取卖家、合同、付款和物流资料，核对钢制螺旋桩与“眼镜柜”提单之间的真实关系。",
        "第二优先：取得Linktrans帖子的聊天、报价、账号和付款主体，对同名品牌进行法律实体消歧后再查申报价值。",
        "第三优先：复核深圳汇升、四川瀚瑞森、企博网账号和宁波—多伦多报价链条，按企业/账号—口岸—订单/提单—实货建立闭环。",
        "二级观察：东莞锦秀公开页面同时强调如实申报；资料一致时应保留为正常代理样本，不作风险升级。",
    ]
    num_id = base.new_num_id(doc, decimal_abstract)
    for action in actions:
        base.add_list_paragraph(doc, action, num_id)

    doc.add_heading("限制与法律提示", level=2)
    limits = [
        "本报告不纳入执法案例，避免与既有执法信息主题重复。",
        "未公开或未能独立核验的企业、个人、货物和单证信息不得补写、推断或外推。",
        "企业或品牌被帖子提及，只表示存在可供核查的节点，不表示其参与违法违规活动。",
        "后续结论须以工商登记、平台订单、合同付款、出口申报、舱单提单、进口申报、实货和代理关系的交叉核验为准。",
    ]
    num_id = base.new_num_id(doc, bullet_abstract)
    for limit in limits:
        base.add_list_paragraph(doc, limit, num_id)


def audit_docx(path: str) -> None:
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        styles_xml = archive.read("word/styles.xml")
        numbering_xml = archive.read("word/numbering.xml")
        relationships_xml = archive.read("word/_rels/document.xml.rels")

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ET.fromstring(document_xml)
    section = root.find(".//w:sectPr", namespace)
    page_size = section.find("w:pgSz", namespace)
    page_margin = section.find("w:pgMar", namespace)
    assert page_size.get(qn("w:w")) == "12240" and page_size.get(qn("w:h")) == "15840"
    for edge in ("top", "right", "bottom", "left"):
        assert page_margin.get(qn(f"w:{edge}")) == "1440"

    tables = root.findall(".//w:tbl", namespace)
    assert len(tables) == 2, f"Expected 2 tables, found {len(tables)}"
    for table in tables:
        width = table.find("w:tblPr/w:tblW", namespace)
        indent = table.find("w:tblPr/w:tblInd", namespace)
        grid = table.findall("w:tblGrid/w:gridCol", namespace)
        assert width.get(qn("w:w")) == str(USABLE_WIDTH_DXA)
        assert indent.get(qn("w:w")) == str(TABLE_INDENT_DXA)
        assert sum(int(node.get(qn("w:w"))) for node in grid) == USABLE_WIDTH_DXA

    text = "".join(node.text or "" for node in root.findall(".//w:t", namespace))
    assert REPORT_TITLE in text
    assert "290841751001029390" in text
    assert "信息源逐源复核台账" in text
    assert all(audit["name"] in text for audit in SOURCE_AUDIT)
    assert "turn" not in text and "cite" not in text
    assert b"List Compact" in styles_xml
    assert b'w:numFmt w:val="bullet"' in numbering_xml
    assert b'w:numFmt w:val="decimal"' in numbering_xml
    assert relationships_xml.count(b"TargetMode=\"External\"") >= len(SOURCE_AUDIT) + len(REPORT_LEADS)
    print(f"audit=passed tables={len(tables)} sources={len(SOURCE_AUDIT)} leads={len(REPORT_LEADS)} chars={len(text)}")


def build() -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    base.configure_page(section)
    configure_header_footer(section)
    base.configure_styles(doc)
    doc.core_properties.title = REPORT_TITLE
    doc.core_properties.subject = "外贸论坛违规清关与走私风险线索全源复核"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "外贸论坛, 报关清关, 中国实体, 口岸核查, 提单, 风险线索"

    bullet_abstract = base.append_num_definition(doc, "bullet")
    decimal_abstract = base.append_num_definition(doc, "decimal")

    add_title_block(doc)
    doc.add_heading("一、本轮结果", level=1)
    add_callout(
        doc,
        "筛选结果：",
        "逐源复核36个既有来源，并从开放互联网新增3个监测/主体核验来源。新增3组近30天线索；与既有4组近月线索合并后，共形成7组当前核查清单，并增加1个中秋、国庆时令生鲜专项预警。",
    )
    add_summary_table(doc)
    note = doc.add_paragraph(style="Small Note")
    note.add_run("判定原则：公开帖子只作待核线索；DDP、包税、双清、LCL或第三国转运单独出现不构成风险升级。")

    for index, lead in enumerate(REPORT_LEADS, 1):
        paragraph = doc.add_paragraph()
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        add_lead_section(doc, lead, index, bullet_abstract, decimal_abstract)

    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    add_seasonal_alert_section(doc, bullet_abstract, decimal_abstract)

    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    add_source_audit_table(doc)

    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    add_method_and_conclusion(doc, bullet_abstract, decimal_abstract)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text:
                existing_size = run.font.size.pt if run.font.size else None
                existing_color = str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None
                base.set_run_font(run, existing_size, existing_color, run.bold, run.italic)

    doc.save(OUTPUT_DOCX)
    audit_docx(OUTPUT_DOCX)
    print(f"docx={OUTPUT_DOCX}")
    return OUTPUT_DOCX


if __name__ == "__main__":
    build()
