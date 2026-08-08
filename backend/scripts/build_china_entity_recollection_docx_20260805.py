"""Build the polished Word version of the China-entity recollection report."""

from __future__ import annotations

import os
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from china_entity_recollection_data_20260805 import (
    DATE_RANGE_END,
    DATE_RANGE_START,
    LEADS,
    OUTPUT_DOCX,
    OUTPUT_DIR,
    REPORT_TITLE,
    SOURCES,
)


CALIBRI = "Calibri"
CJK_FONT = "Microsoft YaHei"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK_BLUE = "0B2545"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "D1D5DB"
RISK_RED = "9B1C1C"
CAUTION = "7A5A00"
USABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = CALIBRI
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), CALIBRI)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), CALIBRI)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size: float, color: str = "000000", bold: bool = False,
                   italic: bool = False) -> None:
    style.font.name = CALIBRI
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), CALIBRI)
    rfonts.set(qn("w:hAnsi"), CALIBRI)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)


def set_spacing(style, before: float, after: float, line: float) -> None:
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    set_spacing(normal, 0, 6, 1.10)

    title = doc.styles["Title"]
    set_style_font(title, 24, INK_BLUE, bold=True)
    set_spacing(title, 0, 6, 1.0)
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 12, MUTED)
    set_spacing(subtitle, 0, 14, 1.10)
    subtitle.paragraph_format.keep_with_next = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        set_style_font(style, size, color, bold=True)
        set_spacing(style, before, after, 1.0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    custom_styles = {
        "Kicker": (10, BLUE, True, False, 0, 4, 1.0),
        "Metadata": (10.2, MUTED, False, False, 0, 2, 1.0),
        "Source Citation": (9, MUTED, False, False, 4, 6, 1.0),
        "List Compact": (11, "000000", False, False, 0, 8, 1.167),
        "Risk Note": (10.5, INK_BLUE, False, False, 4, 8, 1.10),
        "Small Note": (9, MUTED, False, False, 4, 4, 1.05),
        "Table Body": (9.2, "000000", False, False, 0, 0, 1.05),
    }
    for name, values in custom_styles.items():
        if name in doc.styles:
            style = doc.styles[name]
        else:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        size, color, bold, italic, before, after, line = values
        set_style_font(style, size, color, bold=bold, italic=italic)
        set_spacing(style, before, after, line)


def add_paragraph_shading(paragraph, fill: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shading = ppr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        ppr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def set_paragraph_indents(paragraph, left: int = 120, right: int = 120,
                          first_line: int | None = None) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:right"), str(right))
    if first_line is not None:
        ind.set(qn("w:firstLine"), str(first_line))


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), CALIBRI)
    rfonts.set(qn("w:hAnsi"), CALIBRI)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([rfonts, color_el, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([rpr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    run = paragraph.add_run()
    set_run_font(run, 8.5, MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("外贸论坛违规清关与走私风险线索")
    set_run_font(left, 8.5, MUTED, bold=True)
    right = paragraph.add_run("\t中国实体新一轮采集")
    set_run_font(right, 8.5, MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("公开网络待核线索｜不构成违法认定")
    set_run_font(left, 8.5, MUTED)
    tab = paragraph.add_run("\t第 ")
    set_run_font(tab, 8.5, MUTED)
    add_field(paragraph, " PAGE ")
    middle = paragraph.add_run(" 页 / 共 ")
    set_run_font(middle, 8.5, MUTED)
    add_field(paragraph, " NUMPAGES ", "5")
    end = paragraph.add_run(" 页")
    set_run_font(end, 8.5, MUTED)


def append_num_definition(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")

    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, ind, spacing])
    level.extend([start, num_fmt, level_text, suffix, justification, ppr])

    if kind == "bullet":
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), "Symbol")
        rfonts.set(qn("w:hAnsi"), "Symbol")
        rpr.append(rfonts)
        level.append(rpr)

    abstract.append(level)
    numbering.append(abstract)
    return abstract_id


def new_num_id(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_paragraph(doc: Document, text: str, num_id: int) -> None:
    paragraph = doc.add_paragraph(style="List Compact")
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.extend([level, num])
    ppr.insert(0, numpr)
    run = paragraph.add_run(text)
    set_run_font(run, 11)


def set_cell_margins(cell, top: int = 80, bottom: int = 80,
                     start: int = 120, end: int = 120) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tblw = tblpr.find(qn("w:tblW"))
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    indent = tblpr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tblpr.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        cant_split = trpr.find(qn("w:cantSplit"))
        if cant_split is None:
            trpr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_table_borders(table) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shading = tcpr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tcpr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        trpr.append(header)
    header.set(qn("w:val"), "true")


def add_labeled_paragraph(doc: Document, label: str, text: str,
                          style: str | None = None, label_color: str = INK_BLUE):
    paragraph = doc.add_paragraph(style=style)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, 11 if style != "Small Note" else 9, label_color, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, 11 if style != "Small Note" else 9, MUTED if style == "Small Note" else "000000")
    return paragraph


def add_summary_table(doc: Document) -> None:
    headers = ["序号", "日期 / 等级", "中国实体或账号", "主要风险组合及口岸"]
    widths = [620, 1680, 2700, 4360]
    table = doc.add_table(rows=1, cols=4)
    table.style = None
    for column, text in enumerate(headers):
        cell = table.rows[0].cells[column]
        shade_cell(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.style = doc.styles["Table Body"]
        run = paragraph.add_run(text)
        set_run_font(run, 9.2, INK_BLUE, bold=True)
    set_repeat_header(table.rows[0])

    for index, lead in enumerate(LEADS, 1):
        row = table.add_row()
        values = [
            f"{index:02d}",
            lead["published_at"][:10] + "\n" + lead["risk_level"],
            lead["china_entities"][0],
            "；".join(lead["risk_signals"][:3]) + "。" + lead["china_entities"][3],
        ]
        for column, value in enumerate(values):
            paragraph = row.cells[column].paragraphs[0]
            paragraph.style = doc.styles["Table Body"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column < 2 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, 9.2)
    set_table_geometry(table, widths)
    set_table_borders(table)
    after = doc.add_paragraph(style="Source Citation")
    run = after.add_run("说明：一级表示优先补齐主体与单证；二级表示先核实代理关系和申报一致性。")
    set_run_font(run, 9, MUTED)


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph(style="Kicker")
    kicker.add_run("风险线索专题报告")
    title = doc.add_paragraph(REPORT_TITLE, style="Title")
    title.paragraph_format.keep_with_next = True
    subtitle = doc.add_paragraph(
        "近一个月｜非执法案例｜中国实体优先｜公开网络待核线索",
        style="Subtitle",
    )
    metadata = [
        ("报告主题：", "外贸论坛违规清关与走私风险线索"),
        ("采集窗口：", f"{DATE_RANGE_START}—{DATE_RANGE_END}"),
        ("生成日期：", "2026-08-05"),
        ("筛选口径：", "强风险信号组合 + 中国企业/联系方式/口岸或运输标识 + 同主体去重"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph(style="Metadata")
        label_run = paragraph.add_run(label)
        set_run_font(label_run, 10.2, INK_BLUE, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, 10.2, MUTED)


def add_source_line(doc: Document, lead: dict) -> None:
    source = next(source for source in SOURCES if source["id"] == lead["source_id"])
    paragraph = doc.add_paragraph(style="Source Citation")
    label = paragraph.add_run("原帖来源：")
    set_run_font(label, 9, MUTED, bold=True)
    add_hyperlink(paragraph, source["name"], lead["url"])
    tail = paragraph.add_run("  ｜  " + lead["url"])
    set_run_font(tail, 8.5, MUTED)


def add_lead_section(doc: Document, lead: dict, index: int,
                     bullet_abstract: int, decimal_abstract: int) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run(f"{index}. {lead['title']}")
    add_source_line(doc, lead)
    add_labeled_paragraph(
        doc,
        "发布时间 / 核查等级：",
        f"{lead['published_at'][:10]}  ｜  {lead['risk_level']}",
        style="Metadata",
    )

    doc.add_heading("采集到的主要内容", level=2)
    paragraph = doc.add_paragraph(lead["content_summary"])
    paragraph.paragraph_format.keep_together = True

    doc.add_heading("中国可核查实体与标识", level=2)
    bullet_id = new_num_id(doc, bullet_abstract)
    for value in lead["china_entities"]:
        add_list_paragraph(doc, value, bullet_id)

    doc.add_heading("风险分析", level=2)
    risk_id = new_num_id(doc, bullet_abstract)
    for value in lead["risk_points"]:
        add_list_paragraph(doc, value, risk_id)

    doc.add_heading("下一步核查建议", level=2)
    step_id = new_num_id(doc, decimal_abstract)
    for value in lead["next_steps"]:
        add_list_paragraph(doc, value, step_id)

    note = doc.add_paragraph(style="Risk Note")
    add_paragraph_shading(note, CALLOUT)
    set_paragraph_indents(note, left=160, right=160)
    label = note.add_run("当前缺失字段：")
    set_run_font(label, 10.5, RISK_RED, bold=True)
    body = note.add_run(lead["missing_fields"])
    set_run_font(body, 10.5, INK_BLUE)

    reliability = doc.add_paragraph(style="Small Note")
    label = reliability.add_run("来源可靠性：")
    set_run_font(label, 9, CAUTION, bold=True)
    text = reliability.add_run(lead["source_note"])
    set_run_font(text, 9, MUTED)


def add_final_section(doc: Document, bullet_abstract: int, decimal_abstract: int) -> None:
    doc.add_heading("综合研判与执行顺序", level=1)
    callout = doc.add_paragraph(style="Risk Note")
    add_paragraph_shading(callout, CALLOUT)
    set_paragraph_indents(callout, left=160, right=160)
    label = callout.add_run("结论：")
    set_run_font(label, 10.5, INK_BLUE, bold=True)
    text = callout.add_run(
        "本轮新增4组中国实体线索。汇升与瀚瑞森的风险组合最直接；企博网账号需先反查公司主体；"
        "锦秀物流公开页面同时强调如实申报，应作为申报主体关系核验样本，而非直接升级。"
    )
    set_run_font(text, 10.5, INK_BLUE)

    doc.add_heading("建议执行顺序", level=2)
    actions = [
        "核准汇升、瀚瑞森工商登记、备案资质和公开联系方式，并以电话、QQ、店铺账号关联跨平台广告及常用报关抬头。",
        "围绕盐田、蛇口、前海湾、大铲湾、文锦渡、皇岗和南沙，以SO、车牌、柜号、封条、VGM及品类筛查近一个月报关、入仓、拖车和查验记录。",
        "对企博网账号通过电话、邮箱、网站备案和地址反查中国公司，再核验日本品牌手办、3C/免3C、实际进口抬头和品牌授权。",
        "对锦秀线路核对备案外贸抬头、实际惠州工厂、FORM E、合同发票、货值、收汇和实货；全部一致则保留为合规代理样本。",
    ]
    action_id = new_num_id(doc, decimal_abstract)
    for action in actions:
        add_list_paragraph(doc, action, action_id)

    doc.add_heading("排除与限制", level=2)
    exclusions = [
        "执法案例不纳入本主题，避免与既有执法信息主题重复。",
        "平台目录、企业供求和职业博客可能包含营销、模板复制或历史内容刷新；时间、主体和业务真实性须以平台后台、网页快照、官方登记及原始单证复核。",
        "本轮未采集到可公开核验的具体报关单号、提单号、船名航次或实际集装箱号，均作为缺失字段列明，不进行推测或编造。",
        "企业、个人、船公司、仓库和口岸被公开帖子提及，只表示可供核查的节点，不表示其参与违法违规活动。",
    ]
    exclusion_id = new_num_id(doc, bullet_abstract)
    for exclusion in exclusions:
        add_list_paragraph(doc, exclusion, exclusion_id)


def audit_docx(path: str) -> None:
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        styles_xml = archive.read("word/styles.xml")
        numbering_xml = archive.read("word/numbering.xml")
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ET.fromstring(document_xml)
    section = root.find(".//w:sectPr", namespace)
    page_size = section.find("w:pgSz", namespace)
    page_margin = section.find("w:pgMar", namespace)
    assert page_size.get(qn("w:w")) == "12240" and page_size.get(qn("w:h")) == "15840"
    for edge in ("top", "right", "bottom", "left"):
        assert page_margin.get(qn(f"w:{edge}")) == "1440"

    tables = root.findall(".//w:tbl", namespace)
    assert tables, "Expected summary table"
    for table in tables:
        width = table.find("w:tblPr/w:tblW", namespace)
        indent = table.find("w:tblPr/w:tblInd", namespace)
        grid = table.findall("w:tblGrid/w:gridCol", namespace)
        assert width.get(qn("w:w")) == str(USABLE_WIDTH_DXA)
        assert indent.get(qn("w:w")) == str(TABLE_INDENT_DXA)
        assert sum(int(node.get(qn("w:w"))) for node in grid) == USABLE_WIDTH_DXA
        expected = [int(node.get(qn("w:w"))) for node in grid]
        for row in table.findall("w:tr", namespace):
            actual = [int(node.get(qn("w:w"))) for node in row.findall("w:tc/w:tcPr/w:tcW", namespace)]
            assert actual == expected

    assert b"List Compact" in styles_xml
    assert b"w:numFmt w:val=\"bullet\"" in numbering_xml
    assert b"w:numFmt w:val=\"decimal\"" in numbering_xml
    text = "".join(node.text or "" for node in root.findall(".//w:t", namespace))
    assert "turn" not in text and "cite" not in text
    assert not any(paragraph.startswith(("- ", "* ", "• ")) for paragraph in text.splitlines())
    print(f"audit=passed tables={len(tables)} chars={len(text)}")


def build() -> str:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    configure_page(section)
    configure_header_footer(section)
    configure_styles(doc)
    doc.core_properties.title = REPORT_TITLE
    doc.core_properties.subject = "外贸论坛违规清关与走私风险线索中国实体核查"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "外贸论坛, 报关清关, 中国实体, 风险线索, 口岸核查"

    bullet_abstract = append_num_definition(doc, "bullet")
    decimal_abstract = append_num_definition(doc, "decimal")

    add_title_block(doc)
    doc.add_heading("一、本轮结果", level=1)
    callout = doc.add_paragraph(style="Risk Note")
    add_paragraph_shading(callout, CALLOUT)
    set_paragraph_indents(callout, left=160, right=160)
    label = callout.add_run("筛选结果：")
    set_run_font(label, 10.5, INK_BLUE, bold=True)
    text = callout.add_run(
        "近一个月重新采集后保留4组新增线索。全部具备中国企业全称或可反查中国工商主体的公开联系方式；"
        "执法案例、超期帖子、仅有境外个人信息及只有“双清/包税”的普通广告已排除。"
    )
    set_run_font(text, 10.5, INK_BLUE)
    add_summary_table(doc)
    intro = doc.add_paragraph(style="Small Note")
    intro.add_run(
        "判定原则：公开帖子只作为待核线索；任何主体是否存在违规，须以申报、单证、实货、资金和代理关系核验结果为准。"
    )

    for index, lead in enumerate(LEADS, 1):
        paragraph = doc.add_paragraph()
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        add_lead_section(doc, lead, index, bullet_abstract, decimal_abstract)

    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    add_final_section(doc, bullet_abstract, decimal_abstract)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text:
                existing_size = run.font.size.pt if run.font.size else None
                existing_color = None
                if run.font.color and run.font.color.rgb:
                    existing_color = str(run.font.color.rgb)
                set_run_font(
                    run,
                    existing_size,
                    existing_color,
                    run.bold,
                    run.italic,
                )

    doc.save(OUTPUT_DOCX)
    audit_docx(OUTPUT_DOCX)
    print(f"docx={OUTPUT_DOCX}")
    return OUTPUT_DOCX


if __name__ == "__main__":
    build()

