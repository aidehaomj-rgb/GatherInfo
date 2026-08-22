"""Import the user-supplied TDK China-to-U.S. supply-chain report under Japan.

The supplied trade screenshot is redacted, so this importer stores one
aggregate shipment and keeps the Patriot end-use edge at follow-up status.
"""
from __future__ import annotations

import os
import shutil
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from app.database import Base, SessionLocal, engine
from app.models import (
    SupplyChainCase,
    SupplyChainEntity,
    SupplyChainEvidence,
    SupplyChainInvestigation,
    SupplyChainOpenSourceEvidence,
    SupplyChainReport,
    SupplyChainShipment,
)


ROOT = Path(__file__).resolve().parents[2]
DB_PATH = ROOT / "data" / "gather.db"
ARCHIVE_PATH = ROOT / "data" / "reports" / "supply_chain" / "TDK在华供应链穿透分析_企业名称补全版.docx"
INVESTIGATION_ID = "inv-japan-tdk-xiamen-us-defense-network"
SHIPMENT_ID = "shp-japan-tdk-xiamen-america-202605-202608-summary"
VISIBLE_TRADE_ROWS = [
    {"arrival_date": "2026-08-07", "weight_kg": 120},
    {"arrival_date": "2026-08-01", "weight_kg": 37},
    {"arrival_date": "2026-07-24", "weight_kg": 54},
    {"arrival_date": "2026-07-16", "weight_kg": 161},
    {"arrival_date": "2026-07-03", "weight_kg": 54},
    {"arrival_date": "2026-07-01", "weight_kg": 315},
    {"arrival_date": "2026-06-20", "weight_kg": 44},
    {"arrival_date": "2026-06-13", "weight_kg": 429},
]


def upsert(db, model, row_id: str, **values):
    row = db.get(model, row_id)
    if row is None:
        row = model(id=row_id, **values)
        db.add(row)
        return row
    for key, value in values.items():
        setattr(row, key, value)
    return row


def utc_date(value: str) -> datetime:
    return datetime.fromisoformat(value).replace(tzinfo=timezone.utc)


def archive_report() -> Path:
    source_value = os.getenv("TDK_REPORT_PATH")
    source = Path(source_value) if source_value else ARCHIVE_PATH
    if not source.is_file():
        raise FileNotFoundError(f"TDK report not found: {source}")
    ARCHIVE_PATH.parent.mkdir(parents=True, exist_ok=True)
    if source.resolve() != ARCHIVE_PATH.resolve():
        shutil.copy2(source, ARCHIVE_PATH)
    return ARCHIVE_PATH


def backup_database() -> Path:
    backup_dir = ROOT / "data" / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"gather.before_tdk_import.{datetime.now():%Y%m%d_%H%M%S}.db"
    with sqlite3.connect(DB_PATH) as source, sqlite3.connect(backup) as target:
        source.backup(target)
    return backup


def report_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [row.text.strip() for row in document.paragraphs if row.text.strip()]
    tables = [
        "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        for table in document.tables
    ]
    return "\n\n".join([*paragraphs, *tables])


def entity(db, row_id: str, **values):
    return upsert(db, SupplyChainEntity, row_id, **values)


def open_source(db, row_id: str, **values):
    return upsert(db, SupplyChainOpenSourceEvidence, row_id, **values)


def relation(db, row_id: str, **values):
    return upsert(db, SupplyChainEvidence, row_id, **values)


def main() -> None:
    report_path = archive_report()
    report_url = report_path.as_uri()
    backup = backup_database()
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        entities = [
            entity(
                db, "ent-japan-tdk-corporation", name="TDK Corporation",
                name_zh="TDK株式会社", country="Japan", entity_type="integrator",
                aliases=["TDK Corp.", "TDK株式会社（TDK Corporation）"], parent_id=None,
                defense_roles=["磁性元件集团整合", "美国集团供应网络", "两用物项关注名单主体"],
                source_url=report_url,
                notes="用户报告将TDK株式会社还原为日本母公司；关注名单原件及名单编号仍需独立复核。",
            ),
            entity(
                db, "ent-japan-tdk-narita-plant", name="TDK Corporation Narita Plant",
                name_zh="TDK成田工厂", country="Japan", entity_type="component_supplier",
                aliases=["Narita Plant"], parent_id="ent-japan-tdk-corporation",
                defense_roles=["铁氧体产品原制造节点", "向厦门转移部分制造工序（报告材料）"],
                source_url=report_url, notes="产能转移依据来自用户报告汇总，需补工厂通知原件。",
            ),
            entity(
                db, "ent-china-tdk-xiamen", name="TDK Xiamen Co., Ltd.",
                name_zh="厦门TDK有限公司", country="China", entity_type="china_exporter",
                aliases=["厦门TDK", "厦门TDK有限公司（TDK Xiamen Co., Ltd.）"],
                parent_id="ent-japan-tdk-corporation",
                defense_roles=["铁氧体磁芯制造", "磁性元件中国制造节点", "2026年向美国集团公司供货"],
                source_url=report_url,
                notes="报告称主要生产铁氧体磁芯等产品；10票贸易的主体字段在截图中被打码。",
            ),
            entity(
                db, "ent-china-tdk-dalian-electronics", name="TDK大连电子有限公司",
                name_zh="TDK大连电子有限公司", country="China", entity_type="component_supplier",
                aliases=["TDK Dalian Electronics"], parent_id="ent-japan-tdk-corporation",
                defense_roles=["叠层电感器及射频组件制造", "北方制造节点"],
                source_url=report_url, notes="作为TDK在华布局背景节点；未与Patriot具体批次闭合。",
            ),
            entity(
                db, "ent-china-tdk-zhuhai-ftz", name="TDK（珠海保税区）有限公司",
                name_zh="TDK（珠海保税区）有限公司", country="China", entity_type="component_supplier",
                aliases=["TDK Zhuhai Free Trade Zone"], parent_id="ent-japan-tdk-corporation",
                defense_roles=["压敏电阻及薄膜电容器制造", "南方制造节点"],
                source_url=report_url,
                notes="与TDK（珠海）有限公司必须分开；本调查未将其产品归入厦门磁芯批次。",
            ),
            entity(
                db, "ent-us-tdk-corporation-america", name="TDK Corporation of America",
                name_zh="美国TDK公司", country="United States", entity_type="importer",
                aliases=["TDK Corporation of America"], parent_id="ent-japan-tdk-corporation",
                defense_roles=["美国进口、库存与集团渠道承接"], source_url=report_url,
                notes="报告将其识别为厦门TDK货物的美国进口端；需逐票提单复核。",
            ),
            entity(
                db, "ent-us-tdk-lambda-americas", name="TDK-Lambda Americas Inc.",
                name_zh="TDK-Lambda美国公司", country="United States",
                entity_type="defense_supplier", aliases=["TDK-Lambda Americas"],
                parent_id="ent-japan-tdk-corporation",
                defense_roles=["直流电源", "军用维修方舱电源", "美国空军自动测试站套件"],
                source_url=report_url,
                notes="军方采购线与厦门磁芯贸易线属于集团层交叉，不能自动拼接为同一实物链。",
            ),
            entity(
                db, "ent-us-dla-tdk-network", name="U.S. Defense Logistics Agency",
                name_zh="美国国防后勤局", country="United States",
                entity_type="government_agency", aliases=["DLA"], parent_id=None,
                defense_roles=["供应链图谱与军用电源采购（报告材料）"], source_url=report_url,
                notes="报告提及2024年供应链图谱和采购记录，但未附原件页码或链接。",
            ),
        ]
        db.flush()

        cases = [
            upsert(
                db, SupplyChainCase, "case-japan-tdk-xiamen-us-core-trade",
                country="Japan", title="厦门TDK磁芯直供美国集团公司（2026年10票汇总）",
                procurement_agency="TDK Corporation of America",
                procurement_reference="用户报告/打码贸易截图汇总",
                procurement_date=utc_date("2026-08-07"),
                supplier_entity_id="ent-japan-tdk-corporation",
                product="铁氧体磁芯等磁性元件（完整货描和料号待核）",
                target_program="TDK集团美国供应、库存或再分销网络",
                contract_value=None, currency="USD", source_url=report_url,
                source_excerpt="报告称2026年5月至8月至少10票，合计2421千克、193箱；截图字段部分打码。",
                status="researching",
            ),
            upsert(
                db, SupplyChainCase, "case-japan-tdk-patriot-network-lead",
                country="Japan", title="TDK/厦门TDK与Patriot供应网络图谱关联（待原件核验）",
                procurement_agency="U.S. Defense Logistics Agency",
                procurement_reference="报告所述2024年DLA供应链图谱",
                procurement_date=None, supplier_entity_id="ent-japan-tdk-corporation",
                product="磁性元件/铁氧体磁芯（具体Patriot料号待核）",
                target_program="Patriot防空导弹供应网络",
                contract_value=None, currency="USD", source_url=report_url,
                source_excerpt="报告称DLA图谱把Patriot、TDK株式会社和厦门TDK纳入同一供应网络；尚缺图谱原件、页码和节点定义。",
                status="follow_up",
            ),
            upsert(
                db, SupplyChainCase, "case-japan-tdk-lambda-us-military-power",
                country="Japan", title="TDK-Lambda Americas电源进入美国军方维修测试体系（报告线索）",
                procurement_agency="U.S. Defense Logistics Agency / U.S. Air Force",
                procurement_reference="用户报告所列采购公告摘要",
                procurement_date=None, supplier_entity_id="ent-us-tdk-lambda-americas",
                product="直流电源及自动测试站套件",
                target_program="A***M-***电子设备维修方舱/多用途维修基地自动测试站",
                contract_value=None, currency="USD", source_url=report_url,
                source_excerpt="报告称DLA采购TDK-Lambda直流电源，2026年美国空军又以单一来源采购其套件。",
                status="follow_up",
            ),
        ]
        db.flush()

        shipment = upsert(
            db, SupplyChainShipment, SHIPMENT_ID,
            exporter_name="厦门TDK有限公司（TDK Xiamen Co., Ltd.）",
            exporter_country="China", importer_entity_id="ent-us-tdk-corporation-america",
            importer_name="美国TDK公司（TDK Corporation of America）",
            product="铁氧体磁芯等磁性元件（截图货描部分打码）", hs_code=None,
            shipment_date=utc_date("2026-08-07"), weight_kg=2421.0,
            quantity=10.0, quantity_unit="shipments", origin_country="China",
            destination_country="United States", bill_no=None,
            source_name="用户上传报告/贸易截图汇总", source_url=report_url,
            raw_record={
                "evidence_boundary": "report_aggregate_not_individual_bills",
                "reported_period": "2026-05 to 2026-08",
                "reported_shipments": 10,
                "reported_total_weight_kg": 2421,
                "reported_total_packages": 193,
                "visible_rows": VISIBLE_TRADE_ROWS,
                "visible_rows_weight_kg": sum(row["weight_kg"] for row in VISIBLE_TRADE_ROWS),
                "bill_numbers_redacted": True,
                "importer_and_supplier_fields_redacted": True,
                "product_text_partially_redacted": True,
                "source_report": str(report_path),
            },
        )
        db.flush()

        evidence = [
            relation(
                db, "ev-japan-tdk-xiamen-america-trade", case_id=cases[0].id,
                shipment_id=shipment.id, relation_type="direct_trade_record",
                evidence_grade="B", score=68, status="follow_up",
                reasoning="报告和截图支持中国原产磁芯类货物输美及日期重量汇总，但提单号、主体字段和完整货描被打码，需商业数据库底稿逐票复核。",
                verified_facts=["报告汇总10票", "合计2421千克、193箱", "截图可见8条日期重量", "原产国标注中国"],
                model_review={"review": "manual import", "boundary": "redacted aggregate trade screenshot"},
                is_reportable=False,
            ),
            relation(
                db, "ev-japan-tdk-xiamen-patriot-candidate", case_id=cases[1].id,
                shipment_id=shipment.id, relation_type="candidate_component",
                evidence_grade="C", score=42, status="follow_up",
                reasoning="TDK/厦门TDK被报告称出现在Patriot供应网络图谱，且存在近期集团内输美贸易；但没有图谱原件、精确料号、BOM、订单或最终用途文件证明该批磁芯进入Patriot。",
                verified_facts=["集团与工厂身份已在报告中还原", "近期输美贸易与图谱关联为两类独立线索", "具体装机关系未闭合"],
                model_review={"review": "manual import", "boundary": "no batch-to-Patriot closure"},
                is_reportable=False,
            ),
            relation(
                db, "ev-japan-tdk-group-us-military-context", case_id=cases[2].id,
                shipment_id=shipment.id, relation_type="corporate_group_context",
                evidence_grade="C", score=30, status="follow_up",
                reasoning="报告支持TDK-Lambda Americas存在独立美国军方采购，但它与厦门TDK磁芯贸易属于不同业务线，只能作为集团军方业务背景，不能合并为同一批次供应关系。",
                verified_facts=["TDK-Lambda军方采购为独立业务线", "厦门磁芯批次未指向TDK-Lambda", "不得以集团关系替代料号和最终用途证据"],
                model_review={"review": "manual import", "boundary": "corporate-level overlap only"},
                is_reportable=False,
            ),
        ]

        open_sources = [
            open_source(
                db, "ose-japan-tdk-report-trade-screenshot", case_id=cases[0].id,
                title="用户报告及贸易截图：厦门TDK磁芯类货物输美汇总",
                source_type="embedded_trade_screenshot", source_publisher="用户上传报告",
                source_url=report_url,
                source_excerpt="报告称2026年5月至8月至少10票、2421千克、193箱；截图可见8行日期、原产国中国及重量，关键主体和提单字段被打码。",
                verified_facts=["2026-06-13至08-07截图可见8行", "可见8行重量合计1214千克", "报告总口径为10票2421千克193箱"],
                evidence_grade="B", status="follow_up",
                limitations=["未提供逐票底稿", "提单号和主体字段被打码", "完整货描和料号不可见"],
            ),
            open_source(
                db, "ose-japan-tdk-report-entity-resolution", case_id=cases[0].id,
                title="企业名称补全：TDK株式会社、厦门TDK及美国集团公司",
                source_type="user_report", source_publisher="用户上传报告",
                source_url=report_url,
                source_excerpt="补充材料将1公司还原为TDK株式会社、2公司还原为厦门TDK有限公司、4子公司还原为TDK Corporation of America。",
                verified_facts=["TDK集团控制关系", "厦门TDK中国制造节点", "美国TDK进口承接节点"],
                evidence_grade="B", status="follow_up",
                limitations=["贸易截图中的主体字段打码", "仍需企业注册和逐票提单交叉确认"],
            ),
            open_source(
                db, "ose-japan-tdk-report-narita-transfer", case_id=cases[0].id,
                title="报告材料：TDK成田工厂铁氧体制造工序向厦门转移",
                source_type="user_report", source_publisher="用户上传报告",
                source_url=report_url,
                source_excerpt="报告称公开工厂通知显示铁氧体制造由成田分阶段转至厦门，粉料工序另涉及稻仓工厂。",
                verified_facts=["报告记录日本产能向厦门转移"], evidence_grade="C",
                status="follow_up", limitations=["未附工厂通知原件和具体日期"],
            ),
            open_source(
                db, "ose-japan-tdk-report-patriot-network", case_id=cases[1].id,
                title="报告线索：DLA 2024年供应链图谱中的TDK/厦门TDK/Patriot关联",
                source_type="user_report", source_publisher="用户上传报告",
                source_url=report_url,
                source_excerpt="报告称图谱将TDK株式会社列在DLA授权供应商一栏，并将厦门TDK列在中国供应商一栏。",
                verified_facts=["报告声称三者处于同一供应网络"], evidence_grade="C",
                status="follow_up", limitations=["缺少图谱原件、页码、图例和供应商节点定义", "不能证明具体磁芯批次装入Patriot"],
            ),
            open_source(
                db, "ose-japan-tdk-report-lambda-military-procurement", case_id=cases[2].id,
                title="报告线索：TDK-Lambda Americas美国军方电源采购",
                source_type="user_report", source_publisher="用户上传报告",
                source_url=report_url,
                source_excerpt="报告称DLA采购TDK-Lambda直流电源用于维修方舱，2026年美国空军单一来源采购其测试站套件。",
                verified_facts=["报告记录TDK集团美国军方采购业务"], evidence_grade="C",
                status="follow_up", limitations=["未附采购公告原始链接和合同号", "与厦门磁芯贸易不是同一已闭合链路"],
            ),
            open_source(
                db, "ose-japan-tdk-evidence-boundary", case_id=cases[1].id,
                title="报告明确边界：Patriot实物装机仍待证",
                source_type="evidence_boundary", source_publisher="用户上传报告",
                source_url=report_url,
                source_excerpt="补充材料明确说明，同一供应网络不等同于具体料号已装入Patriot导弹。",
                verified_facts=["缺料号", "缺BOM", "缺采购订单", "缺最终用途文件"],
                evidence_grade="A", status="verified",
                limitations=["本条用于限定结论，不是Patriot供应关系的正向证明"],
            ),
        ]

        report_content = (
            "# 系统导入说明\n\n"
            "本报告按日本方向归档。贸易事实采用用户材料的聚合口径，不虚构被打码的提单号、主体字段、完整货描或料号。"
            "Patriot图谱和TDK-Lambda军方采购分别作为待补证线索；当前不得表述为厦门磁芯已经进入Patriot导弹。\n\n"
            "# 附件正文与企业名称补全材料\n\n" + report_text(report_path)
        )
        report = upsert(
            db, SupplyChainReport, "scr-japan-tdk-xiamen-us-defense-network",
            country="Japan", title="TDK在华供应链：厦门磁芯输美与Patriot网络关联分析",
            case_ids=[row.id for row in cases], evidence_ids=[row.id for row in evidence],
            model_id="manual-user-report-import", status="completed", content=report_content,
            summary="TDK日本母公司—厦门TDK—美国TDK贸易方向形成报告级线索；10票2421千克/193箱需逐票底稿复核，Patriot装机和TDK-Lambda业务线均未与该批磁芯闭合。",
            error_log=None, generated_at=datetime.now(timezone.utc),
        )

        investigation = upsert(
            db, SupplyChainInvestigation, INVESTIGATION_ID, country="Japan",
            name="TDK株式会社—厦门TDK—美国集团/军工网络供应链",
            description=(
                "用户报告还原TDK株式会社、厦门TDK有限公司、TDK Corporation of America和TDK-Lambda Americas等主体。"
                "报告称2026年5月至8月厦门TDK至少10票磁芯类货物进入美国集团公司，合计2421千克、193箱；"
                "贸易截图关键字段打码，当前以聚合记录保存。报告另称DLA供应链图谱把TDK/厦门TDK与Patriot置于同一网络，"
                "并列出TDK-Lambda美国军方采购。因缺图谱原件、料号、BOM、订单和最终用途文件，调查保持researching。"
            ),
            status="researching",
            entity_ids=[row.id for row in entities], case_ids=[row.id for row in cases],
            shipment_ids=[shipment.id], evidence_ids=[row.id for row in evidence],
            open_source_evidence_ids=[row.id for row in open_sources], report_ids=[report.id],
            updated_at=datetime.now(timezone.utc),
        )
        db.commit()
        print(
            f"imported={investigation.id} entities={len(entities)} cases={len(cases)} "
            f"shipments=1 evidence={len(evidence)} open_sources={len(open_sources)}"
        )
        print(f"report={report_path}")
        print(f"backup={backup}")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
