"""Build the Codex-authored report for the traceable enforcement case batch."""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.database import SessionLocal  # noqa: E402
from app.models import CollectedItem, Report, SystemConfig, Topic  # noqa: E402
from app.report_export import _write_docx, export_report  # noqa: E402


REPORT_ID = "rpt-codex-enforcement-traceable-20260808"
TOPIC_ID = "weekly-enforcement-intelligence"
RUN_ID = "run-5a98ab9a0a0b"
TITLE = "执法信息采集 综合分析报告（可追溯涉华案例专报）"
ITEM_IDS = [
    "codex-us-bis-bosch-huawei-20260617",
    "codex-us-container-makers-antitrust-20260519",
    "codex-in-supreme-shaoxing-airfly-12-containers-20260415",
    "codex-us-bis-coastal-pva-smic-20260414",
    "codex-in-kash-big-tex-container-20260409",
    "codex-us-shandong-fentanyl-entities-20260325",
    "codex-in-mnr-two-containers-scn-20260313",
    "codex-us-bis-teledyne-flir-shanghai-20260226",
    "codex-in-shri-ji-tianjin-jlhy-20260217",
    "codex-us-bis-applied-materials-smic-20260212",
]

SUMMARY = (
    "本报告复核10起境外官方执法案例，其中美国6起、印度4起。案件覆盖海关申报与归类、"
    "出口管制、毒品前体跨境追诉和国际贸易竞争执法。4起印度海关案件披露16个货柜号、"
    "16份报关单、2个提单号和1个发票号；9起案件直接点名中国大陆企业、收货实体或关联公司。"
    "当前7起已形成行政裁定或执法和解，3起仍处拟处罚通知或刑事起诉阶段。建议优先围绕"
    "涉案境内企业、货柜与提单开展申报回溯、同链路扩线和程序状态跟踪。"
)

CONTENT = r"""## 一、监测概况

本报告对本轮多语种互联网检索形成的10起可追溯境外执法案例进行集中复核。全部案例均回溯至境外海关、美国商务部工业与安全局或美国司法部等官方材料，发布时间覆盖2026年2月12日至6月17日。其中，美国案例6起、印度案例4起；按案件类型划分，出口管制执法4起、海关申报与归类执法4起、毒品前体跨境刑事追诉1起、国际贸易竞争执法1起。

本批案例的突出特点是识别信息较为完整。4起印度海关案件共披露16个货柜号、16份报关单、2个提单号和1个商业发票号；9起案件直接点名中国大陆供应商、涉案企业、收货实体或关联公司。程序状态方面，7起已形成行政裁定或执法和解，1起处于拟处罚通知阶段，2起处于刑事起诉阶段。对尚未作出最终裁判的案件，本报告仅陈述官方指控和程序进展，不将指控直接表述为违法事实。

需要说明的是，官方通报发布时间不等同于涉案行为发生时间。部分美国出口管制案件所涉交易发生于更早年度，本报告将其作为2026年公开执法动向和风险样本分析，不作为2026年新发生交易统计。

## 二、主要特点

### （一）境外海关执法由商品层面向“企业、单证、货柜”链式核查延伸

印度海关4起案件均以进口申报、实货查验、实验室检测或价格核验为切入点，问题主要集中在品名、材质、归类、价格和未申报货物。官方裁定不仅披露商品和处置结果，还逐步形成进口商、境外供应商、报关单、提单、商业发票和货柜号相互印证的证据链。Supreme International案一次涉及12个货柜，M N R Enterprises案同时披露2个货柜号、2份报关单和2个提单号，具备进一步开展航线回溯、同批次比对和关联企业扩线的条件。

### （二）涉华关联由“中国来源”向具体企业和最终使用实体深化

本批10案中，9案能够落到具体中国大陆主体，但主体角色并不相同。印度案件中的Big Tex Impex Limited、Tianjin JLHY Import & Export Co. Ltd.和Shaoxing Airfly Import and Export Co. Ltd.属于境外裁定披露的中国供应端；山东两家化工企业及3家集装箱制造企业属于美国刑事案件中被起诉主体；中芯国际、华为关联实体则主要是美国出口管制案件中的中国收货方、最终使用方或关联实体；FLIR Optoelectronic Technology (Shanghai) Co. Ltd.为被处罚企业的中国关联公司。后续处置时应按供应商、进口商、收货人、最终用户、关联企业和被追诉主体分别建档，防止将不同角色混同。

### （三）半导体、传感器和工业软件仍是美国出口管制执法重点

4起美国商务部工业与安全局案件分别涉及离子注入机、半导体制造用PVA刷具、热成像相机及核心组件、MEMS传感器和汽车软件。Applied Materials案显示设备由美国发往韩国组装后再转运至中国大陆，反映中间生产地和关联公司安排并不能切断出口管制审查；Coastal PVA案显示直接发运与经中国经销商转运均可能被纳入交易链核查；Teledyne FLIR案同时涉及中国大陆关联公司记录保存、香港实体清单地址交易及与未公开名称的中国无人机制造商合作；Bosch案则表明传感器和软件等嵌入式产品也可能成为执法对象。

### （四）境外执法工具呈行政、刑事和竞争执法并行态势

本批案例既有海关没收、重新归类核价、赎回放行和行政罚款，也有出口管制行政和解，还有毒品前体、反垄断领域的刑事起诉。美国司法部对山东两家化工企业和6名中国籍人员的案件仍处起诉阶段；对中集集团、上海寰宇物流装备、新华昌集团等集装箱制造企业的反垄断案件亦尚未经法院最终裁判。境外执法信息监测不能只关注海关通报，还应同步跟踪出口管制、司法起诉、竞争执法及后续法院进展。

## 三、重点风险研判

### （一）纺织品及一般工业品申报差异具有重复出现特征

印度4案中有3案涉及纺织品，反复出现申报品名与实验室检测结果不一致、涂层或层压材质认定差异、归类调整和估价争议。此类案件单票金额未必突出，但相同供应商、进口商、商品描述、货运代理或口岸渠道可能形成重复风险。应重点关注中国出口申报品名与境外进口申报品名是否一致，合同、发票、装箱单及材质检测报告能否相互印证，以及同一企业是否存在连续多票相似差异。

### （二）经第三国加工、经销商转运和关联企业交易成为关键审查节点

Applied Materials案的美国—韩国—中国路线、Coastal PVA案经中国经销商转运的安排，以及Teledyne FLIR案中境外关联公司与上海子公司的分工，均表明执法机关正在穿透形式上的中转和关联公司安排，追查实际收货主体、最终用途和许可证义务。对高技术设备、关键零部件、传感器及软件，应将中转国、组装地、经销商、关联公司和最终用户纳入同一交易链审查。

### （三）公开执法文书可转化为境内核查线索，但需严格区分事实层级

货柜号、提单号、报关单号和发票号可用于反向核查运输路径、申报主体、货代关系和同批次交易；境外官方点名的中国企业可用于开展企业画像、历史申报和关联主体比对。但境外拟处罚通知、刑事起诉和媒体转述的证明力不同，境外执法机关的单方指控也不等同于我国监管结论。线索转化时应保留原始文书、程序状态、证据出处和核验记录，形成“发现线索—回溯单证—多源印证—人工研判—分类处置”的闭环。

## 四、下一步处置建议

### （一）对16个货柜和配套单证开展结构化回溯

优先将货柜号、报关单号、提单号、发票号、进口商、供应商、商品描述和处置结果录入线索台账。对能够匹配境内出口申报的数据，核对出口日期、启运港、目的港、运输工具、航次、商品编码、数量重量和申报价格；对同一供应商或进口商的相邻批次开展横向比对，识别是否存在相同品名、相同路线或相同申报模式。

### （二）对被点名中国大陆主体实施分角色核验

对Big Tex Impex Limited、Tianjin JLHY Import & Export Co. Ltd.、Shaoxing Airfly Import and Export Co. Ltd.等供应端主体，重点核查境内注册名称、历史出口记录、商品申报和境外客户关系；对中芯国际、华为关联实体等收货或最终使用主体，重点关注境外执法文件所述交易时间、产品、许可证要求和转运路径；对山东两家化工企业及集装箱制造企业，持续跟踪起诉书、庭审、认罪协议或裁判结果，避免在程序未终结前作确定性判断。

### （三）建立“同企业、同商品、同路线”自动扩线规则

在系统中对企业英文名及别名、货柜号、提单号、报关单号、商品关键词、执法机关和程序状态设置结构化字段。新增案件入库后，自动检索相同企业、相同商品描述、相同港口路线和相邻时间段记录，并按官方原始文书、官方新闻稿、权威媒体和一般网页划分证据等级。对仅有搜索摘要、无可核验发布日期或无法回溯原始材料的信息，进入候选池，不直接作为正式案件入库。

### （四）形成分级推送和动态跟踪机制

对同时具备中国大陆具体企业、精确单证号码和已生效行政裁定的案件，优先形成核查建议；对只有中国来源但供应商未公开的案件，围绕提单、货柜和进口商继续扩线；对拟处罚通知和刑事起诉案件，设置程序状态提醒，持续跟踪最终裁定。重要线索应明确责任部门、核查时限、反馈结果和后续措施，并用处置结果反向校正检索关键词和审核规则。

## 五、综合判断

本批案例表明，境外涉华执法风险已从单一商品申报问题延伸至跨境交易路径、企业关联关系、最终用户、出口许可、竞争秩序和跨境刑事追诉。对海关风险信息情报工作而言，真正有价值的并非案件数量本身，而是能否从公开材料中提取可核验的企业、货柜、提单、报关单、发票、路线和程序状态，并转化为可回溯、可扩线、可反馈的监管线索。

建议以本批10案为样本，继续完善多语种检索、官方文书优先、实体识别、单证抽取和程序状态跟踪能力，逐步形成覆盖“境外发现—境内核验—关联扩线—处置反馈”的工作机制。现阶段应优先核查4起印度海关案件所涉16个货柜及中国供应端，同时持续跟踪3起未终结案件的后续进展。

## 附录：10起可追溯案例

### 1. 美国BIS处罚博世向华为及关联企业无证出口MEMS传感器和汽车软件案

**案例报道：** 据美国商务部工业与安全局2026年6月17日发布的执法通报，Robert Bosch GmbH因向华为及其关联企业出口MEMS传感器和汽车软件过程中违反美国出口管制规定，与BIS达成约3618万美元行政和解，相关交易价值约7237万美元。该案反映出汽车电子、微机电传感器和配套软件向受限实体供货时，产品技术属性、最终用户及许可证义务均可能成为执法审查重点。

**关联要素：** 涉案企业为Robert Bosch GmbH；中国关联对象为华为及其关联企业；涉案产品为MEMS传感器和汽车软件；交易价值约7237万美元，行政和解金额约3618万美元；官方通报未披露集装箱号、提单号等物流单证。程序状态为行政和解已达成。下一步重点核验产品技术参数、具体收货实体、交易时间及许可证适用情况。

**原始来源：** https://www.bis.gov/press-release/robert-bosch-gmbh-bosch-pay-36-million-penalty-bis-violations-pertaining-shipments-huawei

### 2. 美国司法部起诉中集集团、上海寰宇物流装备、新华昌集团等涉嫌集装箱价格垄断

**案例报道：** 据美国司法部反垄断局2026年5月19日发布的起诉信息，4家集装箱制造企业及7名高管被指控在2019年11月至2024年1月期间限制标准干货集装箱产量并操纵价格。被点名的中国大陆企业包括中国国际海运集装箱（集团）股份有限公司、上海寰宇物流装备有限公司和新华昌集团有限公司。该案目前仍处刑事起诉阶段，相关指控尚需经司法程序确认。

**关联要素：** 涉案中国企业为中国国际海运集装箱（集团）股份有限公司、上海寰宇物流装备有限公司和新华昌集团有限公司；涉案产品为标准干货集装箱；涉案期间为2019年11月至2024年1月；官方起诉材料未披露具体集装箱号、提单号和运输工具。程序状态为刑事起诉、尚未最终裁判。下一步重点跟踪庭审、认罪协议或判决，并评估案件对集装箱制造与国际物流行业的影响。

**原始来源：** https://www.justice.gov/opa/pr/four-worlds-largest-container-manufacturing-companies-and-seven-their-executives-indicted

### 3. 印度海关查处Supreme International进口12个中国来源织物货柜申报不实案

**案例报道：** 据印度蒙德拉海关2026年4月15日发布的行政裁定，Supreme International进口的12个中国来源织物货柜存在申报、归类等问题。裁定材料逐票列明货柜号和报关单号，其中货柜MSMU5383845对应的中国供应商为Shaoxing Airfly Import and Export Co. Ltd.。该案具有同一进口商、多票报关和多货柜集中出现申报差异的特征，可作为境内出口端回溯和同批次扩线的重要样本。

**关联要素：** 印度进口商为Supreme International；已披露中国供应商为Shaoxing Airfly Import and Export Co. Ltd.；货柜号为OCGU8035659、MSDU7789879、MSMU8106952、EITU1352740、TGBU6616724、MSCU5259351、EGSU9364524、CAIU4751697、MSMU6724294、MSDU6326910、CXDU2148338、MSMU5383845；对应报关单号为5614052、5743133、5743135、5814473、5917065、5917409、5917853、5916909、5916910、5916913、5916914、5916911。程序状态为行政裁定已作出。下一步重点核查中国供应商、12票申报之间的同批次关联及涉案货物再出口处置情况。

**原始来源：** https://gujaratcustoms.gov.in/juridictional_commissionerate/public/storage/pdfs/IMoywzLFD9vgixx6iNfaIQTJfCWPkYJWARkm5MVW.pdf

### 4. 美国BIS处罚Coastal PVA向中芯国际北京实体无证出口半导体制造用刷具案

**案例报道：** 据美国商务部工业与安全局2026年4月14日发布的执法通报，Coastal PVA Technology Inc.在18次交易中向中国相关实体供应价值约40.0088万美元的半导体制造用PVA刷具，收货或最终使用实体包括中芯国际北京和北方集成电路制造相关实体，部分货物经一家未公开名称的中国经销商转运。Coastal PVA已就相关出口管制违规与BIS达成行政和解。

**关联要素：** 美国供货企业为Coastal PVA Technology Inc.；中国收货或最终使用实体包括中芯国际北京和北方集成电路制造相关实体；另涉及一家名称未公开的中国经销商；涉案产品为半导体制造用PVA刷具；共涉及18次交易，货值约40.0088万美元；官方通报未披露集装箱号、提单号。程序状态为行政和解已达成。下一步重点核验中国经销商身份、产品归类、最终用途及完整交易链。

**原始来源：** https://www.bis.gov/press-release/bis-reaches-administrative-enforcement-settlement-coastal-pva-technology-inc.

### 5. 印度海关查处Big Tex Impex对Kash International聚酯层压织物申报不实案

**案例报道：** 据印度蒙德拉海关2026年4月9日发布的行政裁定，Kash International进口的一票织物经实验室检测，被认定为聚氨酯层压染色聚酯机织物，与进口申报品名及归类不一致。裁定材料将该票货物与Big Tex Impex Limited、货柜EITU9296026及商业发票NS-12相联系，为对照境内出口申报和境外进口申报提供了明确标识。

**关联要素：** 印度进口商为Kash International；材料所列供应商为Big Tex Impex Limited；货柜号为EITU9296026；报关单号为Z-6816390；商业发票号为NS-12；涉案产品经检测为聚氨酯层压染色聚酯机织物。程序状态为行政裁定已作出。下一步重点比对出口申报、商业发票、材质检测结果和境外进口申报的一致性。

**原始来源：** https://gujaratcustoms.gov.in/juridictional_commissionerate/public/storage/pdfs/hh8BnFlZnBz21Wyh5bj1ppxkdeHBqMd6QVhmIYDD.pdf

### 6. 美国联邦大陪审团起诉山东两家化工企业及6名中国籍人员涉芬太尼化学品案

**案例报道：** 据美国司法部2026年3月25日发布的起诉信息，山东两家化工企业及6名中国籍人员被指控参与与芬太尼制造或掺混所用化学品有关的毒品和洗钱共谋，其中3人另被指控企图向墨西哥湾卡特尔提供实质支持。该案目前处于刑事起诉阶段，被告依法推定无罪，相关事实仍需由法院最终认定。

**关联要素：** 中国大陆企业为Shandong Believe Chemical Company Pte Ltd.、Shandong Ranhang Biotechnology Co. Ltd.；涉案人员为6名中国籍人员；关联对象包括墨西哥湾卡特尔；涉案物品为与芬太尼制造或掺混有关的化学品；官方公开稿未披露集装箱号、提单号和运输工具。程序状态为刑事起诉。下一步重点核查化学品具体名称、企业历史贸易记录、境外收货方和资金链。

**原始来源：** https://www.justice.gov/usao-sdoh/pr/grand-jury-charges-additional-chinese-nationals-pharmaceutical-companies-drug

### 7. 印度税收情报局就M N R Enterprises两个中国来源织物货柜发出拟处罚通知

**案例报道：** 据印度税收情报局及蒙德拉海关2026年3月13日公开的拟处罚通知，M N R Enterprises进口的两个中国来源织物货柜被调查存在申报、归类和估价问题。官方材料同时披露了两只货柜、两份报关单和两份提单，可据此从境外进口端反查中国供应商、启运港、承运关系及相邻批次。该案尚处拟处罚通知阶段，不应表述为最终违法裁定。

**关联要素：** 印度进口商为M N R Enterprises；货柜号为KMTU9295467、UETU7364726；报关单号为2962594、3091762；提单号为KMTCNBO8853676、WSZ25060490；货物为中国来源织物。程序状态为拟处罚通知阶段、尚非最终裁定。下一步重点利用提单和货柜号反查中国供应商、启运港、运输工具及相邻批次。

**原始来源：** https://gujaratcustoms.gov.in/juridictional_commissionerate/public/storage/pdfs/Um6cByp1sGVJCX8PKkGYkbX88UoNxIi5k9BVJgyO.pdf

### 8. 美国BIS处罚Teledyne FLIR及其上海子公司热成像设备出口违规案

**案例报道：** 据美国商务部工业与安全局2026年2月26日发布的执法通报，Teledyne FLIR LLC及其关联企业涉及19项出口管制违规，包括从瑞典向中国无证出口热成像相机、核心和套件，上海关联公司未按规定保存记录，向香港实体清单地址交易，以及与一家名称未公开的中国无人机制造商开展合作。相关企业已与BIS达成100万美元行政和解。

**关联要素：** 中国关联公司为FLIR Optoelectronic Technology (Shanghai) Co. Ltd.；另涉及一家名称未公开的中国无人机制造商及香港实体清单地址；涉案产品为热成像相机、核心和套件；涉及19项违规，行政和解金额为100万美元；官方通报未披露集装箱号、提单号。程序状态为行政和解已达成。不得将未公开的无人机制造商推定为任何具体企业。

**原始来源：** https://www.bis.gov/press-release/bis-reaches-administrative-enforcement-settlement-teledyne-flir-llc-its-affiliates-flir-optoelectronic

### 9. 印度海关查处天津JLHY向Shri Ji出口货柜夹带未申报干壁螺钉案

**案例报道：** 据印度蒙德拉海关2026年2月17日发布的行政裁定，Shri Ji进口的一只货柜申报货物为锚固螺母，实货查验发现其中夹带未申报的镀锌和黑色磷化干壁螺钉。裁定材料将该票货物与天津JLHY进出口有限公司及货柜DPWU2049296相联系，具备从境外查获结果回溯境内出口申报的条件。

**关联要素：** 印度进口商为Shri Ji；中国供应商为Tianjin JLHY Import & Export Co. Ltd.；货柜号为DPWU2049296；报关单号为3609905；申报品名为锚固螺母，查获物品为未申报的镀锌和黑色磷化干壁螺钉。程序状态为行政裁定已作出。下一步重点核查境内出口商品描述、装箱单、重量差异及同一供应商其他批次。

**原始来源：** https://gujaratcustoms.gov.in/juridictional_commissionerate/public/storage/pdfs/UANk0HiDroUC0cRQqjULISbuCAhENg8zxJ4UljnX.pdf

### 10. 美国BIS处罚应用材料经韩国向中芯国际多家大陆实体无证转运离子注入机案

**案例报道：** 据美国商务部工业与安全局2026年2月12日发布的执法通报，Applied Materials相关交易涉及54次未经许可再出口和2次未遂再出口，离子注入机货值约1.26亿美元。相关设备先由美国发往韩国组装，再转运至中国大陆，收货实体涉及中芯国际北京、天津、上海、深圳及华南等多家主体。Applied Materials已就相关出口管制违规与BIS达成约2.52亿美元行政和解。

**关联要素：** 美国企业为Applied Materials；中国收货或最终使用实体涉及中芯国际北京、天津、上海、深圳及华南等主体；涉案产品为离子注入机；运输链路为美国发货、韩国组装、中国大陆收货；涉及54次未经许可再出口和2次未遂再出口，货值约1.26亿美元，和解金额约2.52亿美元；官方通报未披露集装箱号和提单号。程序状态为行政和解已达成。下一步重点核查第三国组装企业、中间关联公司、最终用户及许可证义务。

**原始来源：** https://www.bis.gov/press-release/applied-materials-pay-252-million-penalty-bis-illegally-exporting-semiconductor-manufacturing-equipment
"""


def _style_docx(path: str) -> None:
    """Apply the standard_business_brief preset to this report export."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document(path)
    if doc.paragraphs:
        doc.paragraphs[0].text = TITLE
        doc.paragraphs[0].style = doc.styles["Title"]

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    def style_font(style, size: float, color: str, bold: bool = False) -> None:
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    normal = doc.styles["Normal"]
    style_font(normal, 11, "1F2937")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = doc.styles["Title"]
    style_font(title, 22, "0B2545", bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 16, 8),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style_font(style, size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # The generic exporter maps ##/### one level too deep for this report.
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 2":
            paragraph.style = doc.styles["Heading 1"]
        elif paragraph.style.name == "Heading 3":
            paragraph.style = doc.styles["Heading 2"]
        for run in paragraph.runs:
            if run._element.get_or_add_rPr().rFonts is None:
                run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    if doc.paragraphs:
        metadata = doc.paragraphs[0].insert_paragraph_before()
        doc._body._body.remove(metadata._element)
        doc.paragraphs[0]._element.addnext(metadata._element)
        metadata.style = doc.styles["Subtitle"]
        metadata.paragraph_format.space_before = Pt(0)
        metadata.paragraph_format.space_after = Pt(14)
        metadata.paragraph_format.line_spacing = 1.0
        run = metadata.add_run(
            "监测周期：2026年2月12日至6月17日  |  案例数量：10起  |  生成日期：2026年8月8日"
        )
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string("5F6B7A")
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header.paragraphs[0]
    header.text = "执法信息采集  |  可追溯涉华案例专报"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("6B7280")
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    label = footer.add_run("第 ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    suffix = footer.add_run(" 页")
    for run in (label, suffix):
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("6B7280")
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.save(path)


def _style_pdf(path: str) -> None:
    """Render a CJK-safe PDF with the report title and restrained page furniture."""
    import html
    import re

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    regular_font = r"C:\Windows\Fonts\Deng.ttf"
    bold_font = r"C:\Windows\Fonts\Dengb.ttf"
    pdfmetrics.registerFont(TTFont("Deng", regular_font))
    pdfmetrics.registerFont(TTFont("Deng-Bold", bold_font))

    styles = {
        "title": ParagraphStyle(
            "ReportTitle", fontName="Deng-Bold", fontSize=20, leading=28,
            textColor=colors.HexColor("#0B2545"), spaceAfter=8, alignment=TA_LEFT,
        ),
        "metadata": ParagraphStyle(
            "Metadata", fontName="Deng", fontSize=9.5, leading=14,
            textColor=colors.HexColor("#5F6B7A"), spaceAfter=14, alignment=TA_LEFT,
        ),
        "h1": ParagraphStyle(
            "H1", fontName="Deng-Bold", fontSize=15, leading=22,
            textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=7,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "H2", fontName="Deng-Bold", fontSize=12, leading=18,
            textColor=colors.HexColor("#1F4D78"), spaceBefore=8, spaceAfter=5,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "Body", fontName="Deng", fontSize=10.5, leading=17,
            textColor=colors.HexColor("#1F2937"), spaceAfter=6,
            splitLongWords=True, wordWrap="CJK",
        ),
        "footer": ParagraphStyle(
            "Footer", fontName="Deng", fontSize=8.5, leading=10,
            textColor=colors.HexColor("#6B7280"), alignment=TA_RIGHT,
        ),
    }

    doc = SimpleDocTemplate(
        path,
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=22 * mm,
        bottomMargin=20 * mm,
        title=TITLE,
        author="GatherInfo / Codex",
    )
    story = [
        Paragraph(html.escape(TITLE), styles["title"]),
        Paragraph(
            "监测周期：2026年2月12日至6月17日　|　案例数量：10起　|　生成日期：2026年8月8日",
            styles["metadata"],
        ),
    ]
    for line in CONTENT.strip().splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 3))
            continue
        if stripped.startswith("### "):
            if re.match(r"### (3|6|9)\. ", stripped):
                story.append(PageBreak())
            story.append(Paragraph(html.escape(stripped[4:]), styles["h2"]))
        elif stripped.startswith("## "):
            if stripped.startswith("## 附录"):
                story.append(PageBreak())
            story.append(Paragraph(html.escape(stripped[3:]), styles["h1"]))
        elif stripped.startswith("# "):
            story.append(Paragraph(html.escape(stripped[2:]), styles["h1"]))
        else:
            plain = re.sub(r"\*\*(.+?)\*\*", r"\1", stripped)
            plain = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1（\2）", plain)
            source_match = re.match(r"^(.*原始来源：)\s*(https?://\S+)$", plain)
            if source_match:
                lead, source_url = source_match.groups()
                source_text = (
                    f"{html.escape(lead)}"
                    f"<link href=\"{html.escape(source_url.strip(), quote=True)}\" "
                    "color=\"#2E74B5\">官方原始材料</link>"
                )
                story.append(Paragraph(source_text, styles["body"]))
            else:
                story.append(Paragraph(html.escape(plain), styles["body"]))

    def draw_page(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Deng", 8.5)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(22 * mm, A4[1] - 12 * mm, "执法信息采集  |  可追溯涉华案例专报")
        canvas.drawRightString(A4[0] - 22 * mm, 10 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_page, onLaterPages=draw_page)


def _retitle_text_exports(output_files: dict[str, str]) -> None:
    md_path = output_files.get("md")
    if md_path:
        md = Path(md_path).read_text(encoding="utf-8")
        first_break = md.find("\n")
        if first_break >= 0 and md.startswith("# "):
            md = f"# {TITLE}" + md[first_break:]
        Path(md_path).write_text(md, encoding="utf-8")

    html_path = output_files.get("html")
    if html_path:
        html_doc = Path(html_path).read_text(encoding="utf-8")
        html_doc = html_doc.replace(
            "<title>执法信息采集_情报报告_2026-08-08</title>",
            f"<title>{TITLE}</title>",
            1,
        ).replace(
            "<h1>执法信息采集_情报报告_2026-08-08</h1>",
            f"<h1>{TITLE}</h1>",
            1,
        )
        Path(html_path).write_text(html_doc, encoding="utf-8")


def main() -> int:
    db = SessionLocal()
    try:
        topic = db.get(Topic, TOPIC_ID)
        if topic is None:
            raise RuntimeError(f"Topic not found: {TOPIC_ID}")

        existing_ids = {
            row.id for row in db.query(CollectedItem).filter(CollectedItem.id.in_(ITEM_IDS)).all()
        }
        missing = [item_id for item_id in ITEM_IDS if item_id not in existing_ids]
        if missing:
            raise RuntimeError(f"Collected items not found: {missing}")

        report = db.get(Report, REPORT_ID)
        if report is None:
            report = Report(id=REPORT_ID, topic_id=TOPIC_ID, title=TITLE)
            db.add(report)

        report.title = TITLE
        report.report_type = "analytical"
        report.content = CONTENT.strip()
        report.summary = SUMMARY
        report.status = "completed"
        report.model_id = "codex-authored-web-research"
        report.tokens_used = 0
        report.item_count = len(ITEM_IDS)
        report.item_ids = ITEM_IDS
        report.collection_run_id = RUN_ID
        report.date_range_start = datetime(2026, 2, 12, 0, 0, 0)
        report.date_range_end = datetime(2026, 6, 17, 23, 59, 59)
        report.generated_at = datetime.now()
        report.error_log = None
        report.output_files = None
        report.output_dir = None

        system = db.get(SystemConfig, "global")
        output_files = export_report(report, system, topic, formats=["md", "html", "pdf"])
        preferred_docx = Path(report.output_dir) / "执法信息采集_情报报告_2026-08-08.docx"
        fallback_docx = Path(report.output_dir) / "执法信息采集_情报报告_2026-08-08_附录调整版.docx"
        try:
            _write_docx(str(preferred_docx), TITLE, report.content or "")
            output_files["docx"] = str(preferred_docx)
        except PermissionError:
            _write_docx(str(fallback_docx), TITLE, report.content or "")
            output_files["docx"] = str(fallback_docx)
        report.output_files = output_files
        _retitle_text_exports(output_files)
        if output_files.get("docx"):
            _style_docx(output_files["docx"])
        if output_files.get("pdf"):
            _style_pdf(output_files["pdf"])
        db.commit()
        db.refresh(report)
        print(json.dumps({
            "report_id": report.id,
            "title": report.title,
            "model_id": report.model_id,
            "item_count": report.item_count,
            "collection_run_id": report.collection_run_id,
            "output_files": output_files,
        }, ensure_ascii=False, indent=2))
        return 0
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
