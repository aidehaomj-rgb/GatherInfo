from __future__ import annotations

import re
import sys
from datetime import timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app.database import SessionLocal  # noqa: E402
from app.models import CollectedItem, ResearchJob  # noqa: E402


def clean_text(value: str | None) -> str:
    text = value or ""
    text = re.sub(r"!\[[^]]*]\([^)]*\)", "", text)
    text = re.sub(r"\[([^]]+)]\([^)]*\)", r"\1", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"^\s*#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend((color, underline))
    run.append(props)
    node = OxmlElement("w:t")
    node.text = label
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def entity_text(metadata: dict) -> str:
    research = metadata.get("entity_research") or {}
    entities = research.get("entities") or metadata.get("entities") or []
    values = []
    for entity in entities:
        if isinstance(entity, dict):
            name = entity.get("name") or entity.get("value") or entity.get("canonical_name")
            kind = entity.get("type") or entity.get("entity_type")
            if name:
                values.append(f"{kind}: {name}" if kind else str(name))
        elif entity:
            values.append(str(entity))
    return "；".join(dict.fromkeys(values)) or "系统未提取到可稳定复核的具名实体"


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)


def export(job_id: str) -> Path:
    with SessionLocal() as db:
        job = db.get(ResearchJob, job_id)
        if not job:
            raise SystemExit(f"Research job not found: {job_id}")
        by_id = {
            row.id: row
            for row in db.query(CollectedItem).filter(
                CollectedItem.id.in_(job.result_item_ids or [])
            )
        }
        items = [by_id[item_id] for item_id in (job.result_item_ids or []) if item_id in by_id]
        acceptance = job.acceptance_result or {}

    output_dir = ROOT / "reports" / "research"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"最新执法信息采集清单_2026-08-13_{job_id}.docx"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    title = doc.add_heading("最新境外执法信息采集清单", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("采集窗口：2026年8月7日至8月13日｜任务：" + job_id)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("一、采集说明", level=1)
    metrics = acceptance.get("metrics") or {}
    passed = bool(acceptance.get("passed"))
    doc.add_paragraph(
        f"本次按最新多轮检索、涉华配额、案例审核和实体提取逻辑执行5轮采集，"
        f"最终形成{len(items)}条审核后案例。验收结论：{'通过' if passed else '未完全通过'}。"
        "未满足硬指标的项目如实保留，不以低相关或超期内容补足数量。"
    )
    table = doc.add_table(rows=2, cols=6)
    table.style = "Table Grid"
    labels = ["案例数", "强涉华", "总涉华", "非港法域", "实体提取", "官方来源"]
    values = [
        str(metrics.get("items", len(items))),
        f"{metrics.get('strong_china_ratio', 0):.0%}",
        f"{metrics.get('china_ratio', 0):.0%}",
        str(metrics.get("non_hk_jurisdictions", 0)),
        f"{metrics.get('entity_rate', 0):.0%}",
        f"{metrics.get('official_rate', 0):.0%}",
    ]
    for index, label in enumerate(labels):
        set_cell_text(table.rows[0].cells[index], label, True)
        set_cell_text(table.rows[1].cells[index], values[index])

    doc.add_heading("二、案例目录", level=1)
    for index, item in enumerate(items, 1):
        metadata = item.raw_metadata or {}
        review = metadata.get("enforcement_review") or {}
        trans = metadata.get("translation_zh") or {}
        title_zh = clean_text(trans.get("title_zh")) or clean_text(item.title)
        line = doc.add_paragraph(style="List Number")
        line.add_run(title_zh).bold = True
        line.add_run(
            f"（{review.get('jurisdiction') or '法域待核'}；"
            f"{review.get('china_relevance_label') or '关联度待核'}）"
        )

    doc.add_page_break()
    doc.add_heading("三、案例全文", level=1)
    for index, item in enumerate(items, 1):
        metadata = item.raw_metadata or {}
        review = metadata.get("enforcement_review") or {}
        trans = metadata.get("translation_zh") or {}
        title_zh = clean_text(trans.get("title_zh")) or clean_text(item.title)
        content_zh = clean_text(trans.get("content_zh")) or clean_text(
            trans.get("summary_zh") or item.summary or item.content
        )
        doc.add_heading(f"{index}. {title_zh}", level=2)
        facts = [
            f"发布日期：{item.published_at.strftime('%Y-%m-%d') if item.published_at else '待核'}",
            f"法域：{review.get('jurisdiction') or '待核'}",
            f"执法机构：{review.get('authority') or '待核'}",
            f"案件类型：{review.get('case_type') or '待核'}",
            f"涉华等级：{review.get('china_relevance_label') or '待核'}",
        ]
        doc.add_paragraph("｜".join(facts))
        nexus = clean_text(review.get("mainland_nexus_evidence"))
        if nexus:
            doc.add_paragraph("涉华依据：" + nexus)
        doc.add_paragraph("关键实体：" + entity_text(metadata))
        doc.add_paragraph(content_zh)
        source = doc.add_paragraph("原文链接：")
        add_hyperlink(source, item.url, item.url)
        if index != len(items):
            doc.add_paragraph("─" * 36)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("系统自动采集与审核结果｜生成日期：2026-08-13").font.color.rgb = RGBColor(100, 100, 100)
    doc.save(path)
    return path


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: export_research_job_docx.py <job_id>")
    result = export(sys.argv[1])
    print(result)
